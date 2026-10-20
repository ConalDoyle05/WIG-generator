"""The polish step: turns floor shorthand into professional step text, and flags what it can't.

    "Tight w W1/2 P-004"  ->  "Tighten with a 1/2" wrench Bracket (P-004)."

It is deterministic. Every change comes from the glossary (see wig_glossary.py)
or the workbook's Parts sheet, so the same note always gives the same text, and
anything it does not recognise is flagged, never guessed.

Markers you can type in a note:
    HARD    this step is too complex to expand; the Description is left blank
            for you to write. Must be in capitals.
    PN?     a part number belongs here and you'll look it up in the drawings.
            It becomes "[part number needed]" and the step is flagged.
    NOPART  no part is involved (an inspection, a wipe-down), so a missing part
            number is not flagged.

Each process sheet (Assembly, Testing, Packaging) uses these columns:
    A Step #   B Description   C Raw note   D Status   E Flags   F Draft   G Source note
Status is Flagged, Ready (nothing flagged, waiting for your approval), HARD, or
Approved, which only you set. The generator refuses to build a sheet that has
a Status column until every step on it is Approved (or you build a --draft).

Like wig_import, this module takes explicit paths and imports nothing from the
generator.
"""

from __future__ import annotations

import csv
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Set, Tuple

from openpyxl import load_workbook

from wig_glossary import BASE_WORDS, Glossary

SECTIONS = ("Assembly", "Testing", "Packaging")
STATUS_FLAGGED, STATUS_READY, STATUS_HARD, STATUS_APPROVED = "Flagged", "Ready", "HARD", "Approved"
HEADERS = {3: "Raw note", 4: "Status", 5: "Flags", 6: "Draft", 7: "Source note"}
COL_STEP, COL_DESC, COL_RAW, COL_STATUS, COL_FLAGS, COL_DRAFT, COL_SOURCE = 1, 2, 3, 4, 5, 6, 7
TO_RESOLVE_SHEET = "To Resolve"
PART_PLACEHOLDER = "[part number needed]"

_QUOTES = str.maketrans({
    "“": '"', "”": '"', "″": '"', "‘": "'", "’": "'", "′": "'",
    "–": "-", "—": "-", "‑": "-", "‒": "-", "−": "-", " ": " ",
})
LEADING = "([{"
TRAILING = ".,;:!?)]}"
NUMBER = r"\d+(?:\.\d+)?(?:/\d+)?"
NUMBER_RE = re.compile(NUMBER)
BARE_NUMBER_RE = re.compile(NUMBER)
NUMBER_WITH_UNIT_RE = re.compile(rf"({NUMBER})(.+)")
COUNT_FORM_RE = re.compile(r"(?:\d+[xX]|[xX]\d+)")
WORDISH_RE = re.compile(r"[A-Za-z]")
# Units written straight after the number with no space.
TIGHT_UNITS = {'"', "°", "%"}


@dataclass
class Part:
    number: str
    name: str


@dataclass
class PolishResult:
    draft: str
    flags: List[str]
    status: str
    unknown_words: List[str] = field(default_factory=list)
    parts_used: List[str] = field(default_factory=list)
    tools_used: List[str] = field(default_factory=list)


@dataclass
class Token:
    lead: str
    core: str
    trail: str


def normalize_text(text: str) -> str:
    """Straightens smart quotes and dashes and collapses whitespace."""
    return re.sub(r"\s+", " ", (text or "").translate(_QUOTES)).strip()


def normalize_part_number(value: str) -> str:
    return normalize_text(value).upper().replace(" ", "")


def _split(token: str) -> Token:
    lead = ""
    while token and token[0] in LEADING:
        lead, token = lead + token[0], token[1:]
    trail = ""
    while token and token[-1] in TRAILING:
        trail, token = token[-1] + trail, token[:-1]
    return Token(lead, token, trail)


def _is_placeholder(token: Token) -> bool:
    return token.core.upper() == "PN" and token.trail.startswith("?")


def build_vocabulary(glossary: Glossary, parts: Iterable[Part], learned: Iterable[str] = ()) -> Set[str]:
    """Every word the polish step treats as correctly spelled."""
    vocab: Set[str] = set(BASE_WORDS) | glossary.words | glossary.count_words | glossary.safety_words
    vocab |= {w.lower() for w in learned}
    sources: List[str] = [p.name for p in parts]
    sources += [t.name for t in glossary.tools.values()] + [t.phrase for t in glossary.tools.values()]
    sources += list(glossary.terms.values()) + list(glossary.units.values())
    for text in sources:
        vocab |= set(_words(text))
    return vocab


def _words(text: str) -> List[str]:
    return [w.lower() for w in re.findall(r"[A-Za-z][A-Za-z'-]*", normalize_text(text))]


def is_known_word(word: str, vocab: Set[str]) -> bool:
    w = word.lower().strip("'-")
    if w.endswith("'s"):
        w = w[:-2]
    if not w or w in vocab:
        return True
    if "-" in w:
        return all(is_known_word(part, vocab) for part in w.split("-") if part)
    for suffix in ("s", "es", "ed", "d", "ing", "ly", "er"):
        if w.endswith(suffix) and len(w) > len(suffix) + 1:
            stem = w[: -len(suffix)]
            if stem in vocab:
                return True
            # tightening -> tighten, stopped -> stop
            if suffix in ("ed", "ing") and len(stem) > 2 and stem[-1] == stem[-2] and stem[:-1] in vocab:
                return True
            if suffix in ("ed", "ing") and stem + "e" in vocab:
                return True
    return False


def _singular(word: str) -> str:
    w = word.lower()
    for suffix in ("es", "s"):
        if w.endswith(suffix) and len(w) > len(suffix) + 2:
            return w[: -len(suffix)]
    return w


def _finish(pieces: List[str]) -> str:
    text = " ".join(piece for piece in pieces if piece)
    text = re.sub(r"\s+([.,;:!?)\]}])", r"\1", text)
    text = re.sub(r"([(\[{])\s+", r"\1", text)
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return ""
    match = re.search(r"[A-Za-z]", text)
    if match and text[: match.start()].strip("([{\"' ") == "":
        i = match.start()
        text = text[:i] + text[i].upper() + text[i + 1:]
    if text[-1] not in ".!?":
        text += "."
    return text


def polish_note(raw: str, glossary: Glossary, parts: List[Part], vocab: Set[str]) -> PolishResult:
    """Expands one shorthand note. Never invents content: only glossary and Parts substitutions."""
    text = normalize_text(raw)
    if not text:
        return PolishResult("", ["Empty note."], STATUS_FLAGGED)

    raw_tokens = text.split(" ")
    if "HARD" in raw_tokens:
        return PolishResult("", ["Marked HARD: write this step's description yourself."], STATUS_HARD)

    flags: List[str] = []
    no_part = any(_split(t).core.upper() == "NOPART" for t in raw_tokens)
    placeholder = any(_is_placeholder(_split(t)) for t in raw_tokens)
    tokens = [_split(t) for t in raw_tokens if _split(t).core.upper() != "NOPART"]

    parts_by_number = {normalize_part_number(p.number): p for p in parts if p.number}
    part_words = {_singular(w) for p in parts for w in _words(p.name)}
    count_words = {_singular(w) for w in glossary.count_words} | part_words
    fmt = glossary.settings["part_format"]

    out: List[str] = []
    unknown: List[str] = []
    parts_used: List[str] = []
    tools_used: List[str] = []
    consumed_codes: List[str] = []   # tool codes and part numbers: their digits are not loose numbers
    mentions_part = False

    def emit(token: Token, value: str) -> None:
        out.append(f"{token.lead}{value}{token.trail}")

    def unit_of(core: str) -> Optional[str]:
        return glossary.units.get(core.lower())

    def joined_number(number: str, unit: str) -> str:
        return f"{number}{unit}" if unit in TIGHT_UNITS else f"{number} {unit}"

    def counts_as_count(token: Optional[Token]) -> bool:
        if token is None:
            return False
        if normalize_part_number(token.core) in parts_by_number:
            return True
        return bool(WORDISH_RE.search(token.core)) and _singular(token.core) in count_words

    i = 0
    while i < len(tokens):
        token = tokens[i]
        core = token.core

        # Longest multi-word term first. Punctuation may only sit at the edges of the phrase.
        matched = False
        for n in range(min(glossary.max_term_words, len(tokens) - i), 1, -1):
            window = tokens[i:i + n]
            if any(t.trail for t in window[:-1]) or any(t.lead for t in window[1:]):
                continue
            key = " ".join(t.core.lower() for t in window)
            if key in glossary.terms:
                out.append(f"{window[0].lead}{glossary.terms[key]}{window[-1].trail}")
                i += n
                matched = True
                break
        if matched:
            continue

        if not core:
            emit(token, "")
            i += 1
            continue

        if _is_placeholder(token):
            # "?" was split off as trailing punctuation; it belongs to the marker.
            out.append(f"{token.lead}{PART_PLACEHOLDER}{token.trail[1:]}")
            mentions_part = True
            i += 1
            continue

        part = parts_by_number.get(normalize_part_number(core))
        if part is not None:
            # Already written out as "Bracket (P-004)": leave it, so polishing finished text changes nothing.
            already = (token.lead.endswith("(") and token.trail.startswith(")")
                       and " ".join(out).lower().endswith(part.name.lower()))
            emit(token, part.number if already else fmt.format(name=part.name, number=part.number))
            parts_used.append(part.number)
            consumed_codes.append(core)
            mentions_part = True
            i += 1
            continue

        tool = glossary.tools.get(core.lower())
        if tool is not None:
            emit(token, tool.phrase)
            tools_used.append(tool.code)
            consumed_codes.append(core)
            i += 1
            continue

        # A glossary term wins over the number rules: if it swallows a number, the
        # guard at the end flags it rather than the term being silently ignored.
        if core.lower() in glossary.terms:
            emit(token, glossary.terms[core.lower()])
            i += 1
            continue

        nxt = tokens[i + 1] if i + 1 < len(tokens) else None
        if BARE_NUMBER_RE.fullmatch(core):
            if nxt is not None and not token.trail and not nxt.lead and unit_of(nxt.core):
                out.append(f"{token.lead}{joined_number(core, unit_of(nxt.core))}{nxt.trail}")
                i += 2
                continue
            if not counts_as_count(nxt):
                flags.append(f"Number without a unit: {core}. Add the unit to the note (e.g. {core} N·m or {core} in).")
            emit(token, core)
            i += 1
            continue
        with_unit = NUMBER_WITH_UNIT_RE.fullmatch(core)
        if with_unit and unit_of(with_unit.group(2)):
            emit(token, joined_number(with_unit.group(1), unit_of(with_unit.group(2))))
            i += 1
            continue
        if COUNT_FORM_RE.fullmatch(core):
            emit(token, core.lower())
            i += 1
            continue

        if unit_of(core) and not WORDISH_RE.search(core):
            emit(token, unit_of(core))
            i += 1
            continue

        if glossary.part_number_re.fullmatch(core):
            flags.append(f"Unknown part number: {core} is not on the Parts sheet. Check the number, or add the part.")
            mentions_part = True
            emit(token, core)
            i += 1
            continue

        if WORDISH_RE.search(core):
            if _singular(core) in part_words or _singular(core) in {_singular(w) for w in glossary.count_words}:
                mentions_part = True
            if not is_known_word(core, vocab):
                unknown.append(core)
        emit(token, core)
        i += 1

    if unknown:
        listed = ", ".join(dict.fromkeys(unknown))
        flags.append(f"Unrecognised word(s): {listed}. Fix the typo, or add the word to the glossary.")

    if placeholder:
        flags.append("Part number needed (PN?): look it up in the drawings and replace PN? in the note.")
    elif not parts_used and not no_part:
        mode = glossary.settings["require_part_number"]
        if mode == "always" or (mode == "when_part_mentioned" and mentions_part):
            if not any(f.startswith("Unknown part number") for f in flags):
                flags.append("No part number: add it from the drawings, or type NOPART if no part is involved.")

    lowered_words = set(_words(text))
    hits = sorted(w for w in glossary.safety_words if w in lowered_words)
    if hits:
        flags.append(f"Safety word(s): {', '.join(hits)}. Check this step carefully.")

    draft = _finish(out)

    # Guard on the engine itself: every loose number typed must come through unchanged.
    # A term that swallows a number trips this on purpose.
    loose = " ".join(t.core for t in tokens if t.core not in consumed_codes)
    missing = [n for n in NUMBER_RE.findall(loose) if n not in draft]
    if missing:
        flags.append(f"Internal check: number(s) {', '.join(missing)} from the note are missing from the draft.")

    return PolishResult(
        draft=draft,
        flags=flags,
        status=STATUS_FLAGGED if flags else STATUS_READY,
        unknown_words=list(dict.fromkeys(unknown)),
        parts_used=parts_used,
        tools_used=tools_used,
    )


# -----------------------------------------------------------------------------
# Workbook pass
# -----------------------------------------------------------------------------

@dataclass
class PolishReport:
    counts: Dict[str, int] = field(default_factory=lambda: {s: 0 for s in (STATUS_READY, STATUS_FLAGGED, STATUS_HARD, STATUS_APPROVED)})
    to_resolve: List[Tuple[str, str, str, str, str]] = field(default_factory=list)
    unknown_words: Dict[str, int] = field(default_factory=dict)
    kept_edits: int = 0
    reopened: int = 0
    learned: int = 0
    sections: List[str] = field(default_factory=list)
    backup_path: Optional[Path] = None

    @property
    def total(self) -> int:
        return sum(self.counts.values())


def _text(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text


def read_parts(wb) -> List[Part]:
    """Part # in column A and Part Name in column D, as the generator reads them."""
    if "Parts" not in wb.sheetnames:
        return []
    ws = wb["Parts"]
    parts = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        number = _text(row[0] if len(row) > 0 else None)
        name = _text(row[3] if len(row) > 3 else None)
        if number:
            parts.append(Part(number, name or number))
    return parts


def _read_learned(path: Optional[Path]) -> Set[str]:
    if path is None or not path.exists():
        return set()
    return {line.strip().lower() for line in path.read_text(encoding="utf-8").splitlines() if line.strip()}


def _write_learned(path: Optional[Path], words: Set[str]) -> None:
    if path is None:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(sorted(words)) + "\n", encoding="utf-8")


def update_to_add(path: Optional[Path], unknown: Dict[str, Tuple[int, str]], source: str) -> None:
    """Keeps a running list of unrecognised words, most frequent first."""
    if path is None or not unknown:
        return
    rows: Dict[str, List[str]] = {}
    if path.exists():
        with path.open(newline="", encoding="utf-8") as file:
            for row in csv.DictReader(file):
                if row.get("word"):
                    rows[row["word"]] = [row.get("count", "0"), row.get("example", ""), row.get("last_seen", "")]
    for word, (count, example) in unknown.items():
        key = word.lower()
        previous = int(rows.get(key, ["0"])[0] or 0)
        rows[key] = [str(previous + count), example, source]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as file:
        writer = csv.writer(file)
        writer.writerow(["word", "count", "example", "last_seen"])
        for word, (count, example, seen) in sorted(rows.items(), key=lambda item: (-int(item[1][0]), item[0])):
            writer.writerow([word, count, example, seen])


def polish_workbook(
    workbook_path: Path,
    glossary: Glossary,
    *,
    learned_words_path: Optional[Path] = None,
    to_add_path: Optional[Path] = None,
    sections: Iterable[str] = SECTIONS,
) -> PolishReport:
    """Polishes every process sheet in one workbook and writes the results back.

    Approved rows whose note hasn't changed are left alone. A Description you
    edited by hand is never overwritten: the new draft goes in the Draft column
    and the row is flagged instead.
    """
    workbook_path = Path(workbook_path)
    if not workbook_path.exists():
        raise ValueError(f"No workbook at {workbook_path}. Import the steps first, or create the workbook.")
    wb = load_workbook(workbook_path)
    parts = read_parts(wb)
    report = PolishReport()

    # Words from steps you have approved are correct by definition: learn them.
    learned = _read_learned(learned_words_path)
    before = len(learned)
    for name in sections:
        if name not in wb.sheetnames:
            continue
        for row in wb[name].iter_rows(min_row=2, values_only=True):
            status = _text(row[COL_STATUS - 1] if len(row) >= COL_STATUS else None)
            desc = _text(row[COL_DESC - 1] if len(row) >= COL_DESC else None)
            if status.lower() == STATUS_APPROVED.lower() and desc:
                learned |= set(_words(desc))
    report.learned = len(learned) - before
    vocab = build_vocabulary(glossary, parts, learned)
    unknown: Dict[str, Tuple[int, str]] = {}

    for name in sections:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        has_steps = False
        for column, header in HEADERS.items():
            if not _text(ws.cell(row=1, column=column).value):
                ws.cell(row=1, column=column).value = header
        for r in range(2, ws.max_row + 1):
            cell = lambda c: ws.cell(row=r, column=c)
            step, desc, raw = _text(cell(COL_STEP).value), _text(cell(COL_DESC).value), _text(cell(COL_RAW).value)
            status, last_draft, source = _text(cell(COL_STATUS).value), _text(cell(COL_DRAFT).value), _text(cell(COL_SOURCE).value)
            if not (step or desc or raw):
                continue
            if not raw and desc:
                # Typed straight into Description with no Raw note: treat that text as the note.
                raw = desc
                cell(COL_RAW).value = raw
            if not raw:
                continue
            has_steps = True

            reopened = False
            if status.lower() == STATUS_APPROVED.lower():
                if source == raw:
                    report.counts[STATUS_APPROVED] += 1
                    continue
                reopened = True
                report.reopened += 1

            result = polish_note(raw, glossary, parts, vocab)
            flags = list(result.flags)
            if reopened:
                flags.insert(0, "Note changed after approval: review again.")
            edited_by_hand = bool(desc) and desc not in (last_draft, raw, source)
            if edited_by_hand:
                if result.status != STATUS_HARD and result.draft != desc:
                    flags.append("Description was edited by hand, so it was kept. The new draft is in the Draft column.")
                report.kept_edits += 1
            else:
                cell(COL_DESC).value = result.draft or None

            new_status = result.status
            if new_status == STATUS_READY and flags:
                new_status = STATUS_FLAGGED
            cell(COL_STATUS).value = new_status
            cell(COL_FLAGS).value = " ".join(flags) or None
            cell(COL_DRAFT).value = result.draft or None
            cell(COL_SOURCE).value = raw
            report.counts[new_status] += 1
            if new_status in (STATUS_FLAGGED, STATUS_HARD):
                report.to_resolve.append((name, int(step) if step.isdigit() else step, raw, new_status, " ".join(flags)))
            for word in result.unknown_words:
                count, example = unknown.get(word.lower(), (0, raw))
                unknown[word.lower()] = (count + 1, example)
        if has_steps:
            report.sections.append(name)

    if TO_RESOLVE_SHEET in wb.sheetnames:
        del wb[TO_RESOLVE_SHEET]
    ws = wb.create_sheet(TO_RESOLVE_SHEET)
    ws.append(["Section", "Step #", "Raw note", "Status", "What to do"])
    for row in report.to_resolve:
        ws.append(list(row))
    if not report.to_resolve:
        ws.append(["Nothing to resolve."])
    ws.column_dimensions["C"].width = 50
    ws.column_dimensions["E"].width = 90

    report.backup_path = workbook_path.with_name(workbook_path.name + ".bak")
    shutil.copy2(workbook_path, report.backup_path)
    wb.save(workbook_path)

    _write_learned(learned_words_path, learned)
    report.unknown_words = {word: count for word, (count, _example) in unknown.items()}
    update_to_add(to_add_path, unknown, workbook_path.stem)
    return report
