"""Tests for photo handling: EXIF rotation, shrinking, HEIC, and leaving good photos alone."""

from __future__ import annotations

import hashlib
import io
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from PIL import Image, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

import generate_work_instruction as wig  # noqa: E402
import wig_images  # noqa: E402
from wig_scaffold import SAMPLE_ASSEMBLY, scaffold_project  # noqa: E402


def two_tone(size=(400, 200)) -> Image.Image:
    """Red left half, blue right half, so orientation is visible in the pixels."""
    image = Image.new("RGB", size, (0, 0, 230))
    image.paste((230, 0, 0), (0, 0, size[0] // 2, size[1]))
    return image


def save(image: Image.Image, path: Path, fmt: str, orientation=None) -> Path:
    kwargs = {}
    if orientation:
        exif = Image.Exif()
        exif[wig_images.EXIF_ORIENTATION] = orientation
        kwargs["exif"] = exif
    image.save(path, fmt, **kwargs)
    return path


def opened(source) -> Image.Image:
    return Image.open(source if isinstance(source, str) else io.BytesIO(source.getvalue()))


class PrepareTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.dir = Path(tmp.name)

    def test_a_good_photo_is_embedded_untouched(self):
        for fmt, name in (("PNG", "a.png"), ("JPEG", "b.jpg")):
            path = save(two_tone(), self.dir / name, fmt)
            source, size = wig_images.prepare_for_slide(path)
            self.assertEqual(source, str(path))
            self.assertEqual(size, (400, 200))

    def test_exif_rotated_photo_is_turned_upright(self):
        path = save(two_tone(), self.dir / "phone.jpg", "JPEG", orientation=6)
        source, size = wig_images.prepare_for_slide(path)
        self.assertIsInstance(source, io.BytesIO)
        self.assertEqual(size, (200, 400))
        expected = ImageOps.exif_transpose(Image.open(path)).convert("RGB")
        got = opened(source).convert("RGB")
        self.assertEqual(got.size, expected.size)
        for point in ((50, 50), (150, 350)):
            self.assertTrue(all(abs(a - b) < 40 for a, b in zip(got.getpixel(point), expected.getpixel(point))))

    def test_large_photos_are_shrunk_and_keep_their_format(self):
        png = save(two_tone((4000, 3000)), self.dir / "big.png", "PNG")
        jpg = save(two_tone((4000, 3000)), self.dir / "big.jpg", "JPEG")
        for path, fmt in ((png, "PNG"), (jpg, "JPEG")):
            source, size = wig_images.prepare_for_slide(path)
            self.assertEqual(max(size), wig_images.SLIDE_MAX_PX)
            self.assertEqual(opened(source).format, fmt)

    def test_other_formats_are_converted(self):
        path = save(two_tone(), self.dir / "scan.bmp", "BMP")
        source, _ = wig_images.prepare_for_slide(path)
        self.assertEqual(opened(source).format, "JPEG")

    @unittest.skipUnless(wig_images.HEIC_SUPPORTED, "pillow-heif not installed")
    def test_heic_photos_become_jpeg(self):
        path = save(two_tone(), self.dir / "IMG_0001.HEIC", "HEIF")
        source, size = wig_images.prepare_for_slide(path)
        self.assertEqual(opened(source).format, "JPEG")
        self.assertEqual(size, (400, 200))

    def test_heic_without_pillow_heif_says_what_to_install(self):
        path = self.dir / "IMG_0002.heic"
        path.write_bytes(b"\x00\x00\x00\x18ftypheic" + b"\x00" * 64)
        with mock.patch.object(wig_images, "HEIC_SUPPORTED", False):
            with self.assertRaises(ValueError) as caught:
                wig_images.prepare_for_slide(path)
        self.assertIn("pip install pillow-heif", str(caught.exception))

    def test_unreadable_file_is_a_clear_error(self):
        path = self.dir / "broken.jpg"
        path.write_bytes(b"not an image")
        with self.assertRaises(ValueError) as caught:
            wig_images.prepare_for_slide(path)
        self.assertIn("could not be read", str(caught.exception))


class BuildTests(unittest.TestCase):
    """Hand-added photos in a real build: upright, small, and the originals left alone."""

    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        saved = {n: getattr(wig, n) for n in ("DESIGN_FOLDER", "PART_IMAGE_FOLDER", "TEMPLATE_PATH")}
        self.addCleanup(lambda: [setattr(wig, n, v) for n, v in saved.items()])
        self.made = scaffold_project(Path(tmp.name) / "p", sample=True)
        self.steps = self.made["assembly_dir"] / "Assembly"

    def build(self):
        argv = ["--assembly", SAMPLE_ASSEMBLY,
                "--design-folder", str(self.made["design_folder"]),
                "--part-images-folder", str(self.made["part_images_folder"]),
                "--template-path", str(self.made["template_path"]),
                "--initials", "CD", "--welding", "N", "--steps-per-slide", "4", "--initial-release"]
        with mock.patch("sys.stdout", io.StringIO()), mock.patch("sys.stderr", io.StringIO()) as err:
            code = wig.main(argv)
        self.assertEqual(code, 0, err.getvalue())
        deck = self.made["assembly_dir"] / f"{SAMPLE_ASSEMBLY} REV A.pptx"
        return [shape for slide in Presentation(deck).slides for shape in slide.shapes
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]

    def test_rotated_and_huge_photos_are_fixed_in_the_deck_only(self):
        (self.steps / "1.png").unlink()
        (self.steps / "2.png").unlink()
        rotated = save(two_tone((1200, 600)), self.steps / "1.jpg", "JPEG", orientation=6)
        huge = save(two_tone((5000, 2500)), self.steps / "2.jpg", "JPEG")
        before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in (rotated, huge)}
        sizes = sorted(Image.open(io.BytesIO(pic.image.blob)).size for pic in self.build())
        self.assertIn((600, 1200), sizes)                      # upright, not sideways
        self.assertIn((2000, 1000), sizes)                     # shrunk
        self.assertTrue(all(max(s) <= wig_images.SLIDE_MAX_PX for s in sizes))
        self.assertEqual({p: hashlib.sha256(p.read_bytes()).hexdigest() for p in (rotated, huge)}, before)

    def test_rotated_photo_frame_matches_its_upright_shape(self):
        (self.steps / "1.png").unlink()
        save(two_tone((1200, 600)), self.steps / "1.jpg", "JPEG", orientation=6)
        tall = [p for p in self.build() if Image.open(io.BytesIO(p.image.blob)).size == (600, 1200)]
        self.assertEqual(len(tall), 1)
        self.assertGreater(tall[0].height, tall[0].width)

    @unittest.skipUnless(wig_images.HEIC_SUPPORTED, "pillow-heif not installed")
    def test_heic_step_and_part_photos_are_found_and_embedded(self):
        (self.steps / "3.png").unlink()
        save(two_tone(), self.steps / "3.heic", "HEIF")
        (self.made["part_images_folder"] / "P-001.png").unlink()
        save(two_tone(), self.made["part_images_folder"] / "P-001.HEIC", "HEIF")
        formats = {Image.open(io.BytesIO(p.image.blob)).format for p in self.build()}
        self.assertIn("JPEG", formats)


if __name__ == "__main__":
    unittest.main()
