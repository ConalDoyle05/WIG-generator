"""Tests for the "layout" section of --config: moving and resizing what goes on each slide."""

from __future__ import annotations

import io
import json
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

import generate_work_instruction as wig  # noqa: E402
from wig_scaffold import SAMPLE_ASSEMBLY, scaffold_project  # noqa: E402

SIX_UP = {
    "font_size": 9,
    "slots": [
        {"image": [0.30 + 3.15 * (i % 3), 1.00 + 3.10 * (i // 3), 2.95, 2.30],
         "description": [0.30 + 3.15 * (i % 3), 3.35 + 3.10 * (i // 3), 2.95, 0.45]}
        for i in range(6)
    ],
}


class LayoutConfigTests(unittest.TestCase):
    RUNTIME = ("FOOTER_ENABLED", "ALLOW_DRAFT")

    def setUp(self):
        saved = {n: getattr(wig, n) for n in list(wig.CONFIGURABLE_GLOBALS) + list(self.RUNTIME)}
        self.addCleanup(lambda: [setattr(wig, n, v) for n, v in saved.items()])
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)
        self.made = scaffold_project(self.root / "p", sample=True)

    def config(self, layout) -> Path:
        path = self.root / "wig.json"
        path.write_text(json.dumps({"layout": layout}), encoding="utf-8")
        return path

    def build(self, layout, steps="4"):
        argv = ["--assembly", SAMPLE_ASSEMBLY, "--config", str(self.config(layout)),
                "--design-folder", str(self.made["design_folder"]),
                "--part-images-folder", str(self.made["part_images_folder"]),
                "--template-path", str(self.made["template_path"]),
                "--initials", "CD", "--welding", "N", "--steps-per-slide", steps, "--initial-release"]
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = wig.main(argv)
        return code, err.getvalue()

    def deck(self):
        return Presentation(self.made["assembly_dir"] / f"{SAMPLE_ASSEMBLY} REV A.pptx")

    @staticmethod
    def pictures(slide):
        return [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]

    @staticmethod
    def inside(shape, box):
        x, y, w, h = (Inches(v) for v in box)
        return (shape.left >= x - 2 and shape.top >= y - 2
                and shape.left + shape.width <= x + w + 2 and shape.top + shape.height <= y + h + 2)

    def test_cover_photo_box_moves(self):
        box = [0.5, 2.0, 3.0, 3.0]
        code, err = self.build({"title_image_box": box})
        self.assertEqual(code, 0, err)
        cover = self.pictures(self.deck().slides[0])[0]
        self.assertTrue(self.inside(cover, box))

    def test_a_new_six_up_layout_can_be_used(self):
        code, err = self.build({"process_layouts": {**self.default_layouts(), "6": SIX_UP},
                                "step_badges": {"6": {"diameter": 0.36, "font_size": 12}}}, steps="6")
        self.assertEqual(code, 0, err)
        slide = next(s for s in self.deck().slides
                     if any(sh.has_text_frame and sh.text_frame.text.startswith("Assembly Steps") for sh in s.shapes))
        photos = self.pictures(slide)
        self.assertEqual(len(photos), 4)
        for photo, slot in zip(sorted(photos, key=lambda p: (p.top, p.left)), SIX_UP["slots"]):
            self.assertTrue(self.inside(photo, slot["image"]))
        badges = [sh for sh in slide.shapes if sh.name.startswith(wig.STEP_BADGE_NAME)]
        self.assertEqual({round(b.width / 914400, 2) for b in badges}, {0.36})

    def test_replacing_the_layouts_drops_the_old_ones(self):
        code, err = self.build({"process_layouts": {"6": SIX_UP},
                                "step_badges": {"6": {"diameter": 0.36, "font_size": 12}}}, steps="8")
        self.assertEqual(code, 1)
        self.assertIn("--steps-per-slide must be one of: 6", err)

    def test_fewer_part_cards_per_slide_adds_parts_pages(self):
        code, err = self.build({"part_card_positions": [[0.5, 1.0], [0.5, 2.5], [0.5, 4.0]],
                                "part_image_size": [1.2, 1.2], "part_text_size": [4.0, 1.2]})
        self.assertEqual(code, 0, err)
        parts_pages = [s for s in self.deck().slides
                       if any(sh.has_text_frame and "QTY:" in sh.text_frame.text for sh in s.shapes)]
        self.assertEqual(len(parts_pages), 2)        # 5 sample parts, 3 per page
        card = next(sh for sh in parts_pages[0].shapes if sh.has_text_frame and "QTY:" in sh.text_frame.text)
        self.assertEqual((round(card.width / 914400, 2), round(card.height / 914400, 2)), (4.0, 1.2))

    def test_description_limits_apply(self):
        layout = {"description_char_limit": 10, "description_bottom_limit": 7.0, "description_row_gap": 0.1}
        code, err = self.build(layout)
        self.assertEqual(code, 0, err)
        self.assertEqual((wig.STANDARD_DESCRIPTION_CHAR_LIMIT, wig.LONG_DESCRIPTION_BOTTOM_LIMIT,
                          wig.LONG_DESCRIPTION_ROW_GAP), (10, 7.0, 0.1))

    def default_layouts(self):
        return {str(n): {"font_size": font, "slots": [{"image": list(i), "description": list(d)} for i, d in boxes]}
                for n, (boxes, font) in wig.PROCESS_LAYOUTS.items()}

    def test_bad_layouts_are_rejected_with_a_reason_and_change_nothing(self):
        before = {n: getattr(wig, n) for n in wig.CONFIGURABLE_GLOBALS}
        cases = {
            "unknown key": ({"title_box": [1, 1, 1, 1]}, "Unknown layout key"),
            "box with 3 numbers": ({"title_image_box": [1, 1, 1]}, "[left, top, width, height]"),
            "zero width": ({"title_image_box": [1, 1, 0, 1]}, "width must be greater than 0"),
            "negative": ({"process_title_box": [-1, 1, 1, 1]}, "left must be 0 or more"),
            "typo-sized": ({"title_image_box": [1, 1, 500, 1]}, "at most 20"),
            "text not number": ({"part_image_size": ["1", 1]}, "must be a number"),
            "wrong slot count": ({"process_layouts": {"6": {**SIX_UP, "slots": SIX_UP["slots"][:5]}}}, "exactly 6 slot"),
            "no badge for new layout": ({"process_layouts": {"6": SIX_UP}}, "missing: 6"),
            "bad key": ({"process_layouts": {"six": SIX_UP}}, "whole numbers"),
            "bad font": ({"process_layouts": {"6": {**SIX_UP, "font_size": 200}},
                          "step_badges": {"6": {"diameter": 0.3, "font_size": 12}}}, "font size from 6 to 72"),
            "slot missing image": ({"process_layouts": {"2": {"font_size": 10, "slots": [{"description": [0, 0, 1, 1]}] * 2}}},
                                   '"image" and "description"'),
            "too many cards": ({"part_card_positions": [[0, 0]] * 31}, "1 to 30"),
            "char limit": ({"description_char_limit": 0}, "1 to 2000"),
            "boolean": ({"description_bottom_limit": True}, "must be a number"),
        }
        for label, (layout, message) in cases.items():
            with self.subTest(label):
                with self.assertRaises(ValueError) as caught:
                    wig.load_config(self.config(layout))
                self.assertIn(message, str(caught.exception))
                self.assertEqual({n: getattr(wig, n) for n in wig.CONFIGURABLE_GLOBALS}, before)

    def test_a_bad_layout_with_good_slides_applies_neither(self):
        path = self.root / "mixed.json"
        path.write_text(json.dumps({"slides": {"title": 0}, "paths": {"template_path": "x.pptx"},
                                    "layout": {"title_image_box": [1, 1]}}), encoding="utf-8")
        before = wig.TEMPLATE_PATH
        with self.assertRaises(ValueError):
            wig.load_config(path)
        self.assertEqual(wig.TEMPLATE_PATH, before)


if __name__ == "__main__":
    unittest.main()
