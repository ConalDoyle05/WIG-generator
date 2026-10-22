"""Tests for step times: parsing the Time column, adding up ASSEMBLY TIME, and importing times."""

from __future__ import annotations

import datetime as dt
import io
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
sys.path.insert(0, str(HERE))

import generate_work_instruction as wig  # noqa: E402
import wig_import  # noqa: E402
import wig_time  # noqa: E402
from docx_fixture import make_docx, png_bytes  # noqa: E402
from wig_scaffold import SAMPLE_ASSEMBLY, scaffold_project  # noqa: E402


class ParseTests(unittest.TestCase):
    def test_accepted_forms(self):
        cases = {"45 s": 45, "45s": 45, "2 min": 120, "1.5 min": 90, "1m 30s": 90, "1m30s": 90,
                 "1 min 30 sec": 90, "2 minutes and 10 seconds": 130, "1 h 5 min": 3900,
                 "1:30": 90, "1:02:30": 3750, "1 MIN": 60}
        for text, seconds in cases.items():
            with self.subTest(text):
                self.assertEqual(wig_time.parse_duration(text), seconds)

    def test_blank_is_none(self):
        for value in (None, "", "   "):
            self.assertIsNone(wig_time.parse_duration(value))

    def test_bare_numbers_need_a_unit_from_the_header(self):
        for value in (45, "45", 45.0):
            with self.assertRaises(ValueError) as caught:
                wig_time.parse_duration(value)
            self.assertIn("needs a unit", str(caught.exception))
        self.assertEqual(wig_time.parse_duration(45, wig_time.header_unit("Time (s)")), 45)
        self.assertEqual(wig_time.parse_duration("2", wig_time.header_unit("Time (min)")), 120)

    def test_header_units(self):
        self.assertEqual([wig_time.header_unit(h) for h in ("Time", "Time (s)", "Time (sec)", "Time (min)",
                                                            "Time (h)", "Time sec")],
                         [None, 1, 1, 60, 3600, 1])

    def test_rejected_forms(self):
        for text in ("abc", "45 parsecs", "1:75", "fast", "2 min then 3"):
            with self.subTest(text), self.assertRaises(ValueError):
                wig_time.parse_duration(text)

    def test_excel_time_values(self):
        self.assertEqual(wig_time.parse_duration(dt.time(0, 1, 30)), 90)
        self.assertEqual(wig_time.parse_duration(dt.timedelta(minutes=2)), 120)
        with self.assertRaises(ValueError):
            wig_time.parse_duration(dt.datetime(2026, 1, 1))

    def test_format(self):
        self.assertEqual([wig_time.format_duration(s) for s in (45, 720, 750, 3900, 3930, 0.4)],
                         ["45 s", "12 min", "12 min 30 s", "1 h 05 min", "1 h 05 min 30 s", "0 s"])

    def test_duration_text_detection(self):
        self.assertEqual([wig_time.is_duration_text(t) for t in ("45 s", "1:30", "45", "Tighten 2 min", "")],
                         [True, True, False, False, False])


class BuildTests(unittest.TestCase):
    """A Time column on the sample assembly, built through the command line."""

    GLOBALS = ("DESIGN_FOLDER", "PART_IMAGE_FOLDER", "TEMPLATE_PATH", "FOOTER_ENABLED", "ALLOW_DRAFT")

    def setUp(self):
        saved = {n: getattr(wig, n) for n in self.GLOBALS}
        self.addCleanup(lambda: [setattr(wig, n, v) for n, v in saved.items()])
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.made = scaffold_project(Path(tmp.name) / "p", sample=True)
        self.workbook = self.made["assembly_dir"] / f"{SAMPLE_ASSEMBLY}.xlsx"

    def set_times(self, times, header="Time", column=3, assembly_time=None, sheet="Assembly"):
        wb = load_workbook(self.workbook)
        ws = wb[sheet]
        ws.cell(row=1, column=column).value = header
        for row, value in enumerate(times, start=2):
            ws.cell(row=row, column=column).value = value
        wb["Title Info"]["F2"] = assembly_time
        wb.save(self.workbook)

    def run_cli(self, *extra):
        argv = ["--assembly", SAMPLE_ASSEMBLY, "--design-folder", str(self.made["design_folder"]),
                "--part-images-folder", str(self.made["part_images_folder"]),
                "--template-path", str(self.made["template_path"]), *extra]
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = wig.main(argv)
        return code, out.getvalue(), err.getvalue()

    def build(self):
        return self.run_cli("--initials", "CD", "--welding", "N", "--steps-per-slide", "4", "--initial-release")

    def cover_texts(self):
        deck = self.made["assembly_dir"] / f"{SAMPLE_ASSEMBLY} REV A.pptx"
        return [sh.text_frame.text for sh in Presentation(deck).slides[0].shapes if sh.has_text_frame]

    def test_blank_assembly_time_is_added_up_from_the_steps(self):
        self.set_times(["45 s", "1m 30s", "2 min", "1:05"])
        code, out, err = self.build()
        self.assertEqual(code, 0, err)
        self.assertIn("Assembly time: 5 min 20 s", self.cover_texts())
        self.assertIn("Step times: Assembly 5 min 20 s. Total 5 min 20 s.", out)

    def test_time_column_is_found_by_its_header_anywhere(self):
        self.set_times(["30 s"] * 4, column=9)
        code, _, err = self.build()
        self.assertEqual(code, 0, err)
        self.assertIn("Assembly time: 2 min", self.cover_texts())

    def test_testing_times_count_toward_the_total(self):
        self.set_times(["1 min"] * 4)
        wb = load_workbook(self.workbook)
        ws = wb["Testing"]
        ws["C1"] = "Time"
        ws.append([1, "Check the fit", "30 s"])
        wb.save(self.workbook)
        (self.made["assembly_dir"] / "Testing").mkdir()
        (self.made["assembly_dir"] / "Testing" / "1.png").write_bytes(png_bytes())
        code, out, err = self.build()
        self.assertEqual(code, 0, err)
        self.assertIn("Assembly time: 4 min 30 s", self.cover_texts())
        self.assertIn("Step times: Assembly 4 min; Testing 30 s. Total 4 min 30 s.", out)

    def test_missing_step_times_with_a_blank_assembly_time_are_listed(self):
        self.set_times(["45 s", None, "2 min", None])
        code, _, err = self.build()
        self.assertEqual(code, 1)
        self.assertIn("these steps have no time", err)
        self.assertIn("Assembly 2, Assembly 4", err)

    def test_blank_assembly_time_and_no_times_says_both_ways_to_fix_it(self):
        wb = load_workbook(self.workbook)
        wb["Title Info"]["F2"] = None
        wb.save(self.workbook)
        code, _, err = self.build()
        self.assertEqual(code, 1)
        self.assertIn("ASSEMBLY TIME", err)
        self.assertIn("give every step a time", err)

    def test_typed_assembly_time_wins_but_a_mismatch_is_noted(self):
        self.set_times(["1 min"] * 4, assembly_time="5 min")
        code, out, err = self.build()
        self.assertEqual(code, 0, err)
        self.assertIn("Assembly time: 5 min", self.cover_texts())
        self.assertIn("step times add up to 4 min", out)

    def test_matching_typed_time_needs_no_note(self):
        self.set_times(["1 min"] * 4, assembly_time="4 min")
        code, out, _ = self.build()
        self.assertEqual(code, 0)
        self.assertNotIn("add up to", out)

    def test_unreadable_time_is_an_error_naming_the_row(self):
        self.set_times(["45 s", "soon", "2 min", "1 min"], assembly_time="5 min")
        code, _, err = self.build()
        self.assertEqual(code, 1)
        self.assertIn("row 3: the time for step 2 (soon) can't be read", err)

    def test_bare_numbers_follow_the_header_unit(self):
        self.set_times([45, 45, 45, 45], header="Time (s)")
        code, _, err = self.build()
        self.assertEqual(code, 0, err)
        self.assertIn("Assembly time: 3 min", self.cover_texts())
        self.set_times([45, 45, 45, 45], header="Time")
        code, _, err = self.build()
        self.assertEqual(code, 1)
        self.assertIn("needs a unit", err)

    def test_check_reports_times_and_warns_about_excel_clock_values(self):
        self.set_times([dt.time(1, 30), "1 min", "1 min", "1 min"])
        code, out, _ = self.run_cli("--check")
        self.assertEqual(code, 0, out)
        self.assertIn("Step times: Assembly 1 h 33 min. Total 1 h 33 min.", out)
        self.assertIn("Assembly step 1 is timed at 1 h 30 min", out)
        self.assertIn("Excel may have read it as 1 hour 30 minutes", out)


class ImportTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)
        self.assembly = self.root / "12" / "12-345-67"
        self.workbook = self.assembly / "12-345-67.xlsx"

    def run_import(self, rows, **kwargs):
        docx = make_docx(self.root / "page.docx", rows)
        return wig_import.import_onenote(docx, self.assembly, self.workbook, **kwargs)

    def rows(self):
        ws = load_workbook(self.workbook)["Assembly"]
        return [[c.value for c in r] for r in ws.iter_rows(min_row=1, max_col=9)]

    def test_a_time_cell_goes_to_the_time_column_not_the_note(self):
        self.run_import([("Tight w W1/2 P-004", png_bytes(), "45 s"), ("Insp NOPART", png_bytes(), "1m 30s")])
        rows = self.rows()
        self.assertEqual(rows[0][8], "Time")
        self.assertEqual([(r[2], r[8]) for r in rows[1:]], [("Tight w W1/2 P-004", "45 s"), ("Insp NOPART", "1m 30s")])

    def test_a_lone_time_like_note_stays_a_note(self):
        self.run_import([("45 s", png_bytes())])
        self.assertEqual(self.rows()[1][2], "45 s")

    def test_a_typed_time_survives_a_reimport_of_the_same_note(self):
        self.run_import([("Insp NOPART", png_bytes())])
        wb = load_workbook(self.workbook)
        wb["Assembly"].cell(row=2, column=9).value = "2 min"
        wb.save(self.workbook)
        self.run_import([("Insp NOPART", png_bytes())], replace=True)
        self.assertEqual(self.rows()[1][8], "2 min")

    def test_polish_leaves_times_alone(self):
        import copy
        import wig_glossary
        import wig_polish
        self.run_import([("Tight w W1/2 P-004", png_bytes(), "45 s")])
        glossary = wig_glossary.build_glossary(copy.deepcopy(wig_glossary.STARTER_GLOSSARY))
        wig_polish.polish_workbook(self.workbook, glossary)
        self.assertEqual(self.rows()[1][8], "45 s")
        self.assertEqual(self.rows()[0][8], "Time")


if __name__ == "__main__":
    unittest.main()
