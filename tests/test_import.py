"""Tests for the OneNote (.docx) import."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent))

import wig_import  # noqa: E402
from docx_fixture import jpeg_bytes, make_docx, png_bytes  # noqa: E402

RED, GREEN, BLUE = (230, 20, 20), (20, 200, 20), (20, 20, 230)


def dominant(path: Path):
    """Colour of the middle pixel, to prove which photo ended up as which step."""
    with Image.open(path) as image:
        return image.convert("RGB").getpixel((image.width // 2, image.height // 2))


def near(actual, expected, tolerance=40) -> bool:
    return all(abs(a - e) <= tolerance for a, e in zip(actual, expected))


class ImportTestCase(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)
        self.assembly_dir = self.root / "12" / "12-345-67"
        self.workbook = self.assembly_dir / "12-345-67.xlsx"
        self.images = self.assembly_dir / "Assembly"

    def docx(self, rows, **kwargs) -> Path:
        return make_docx(self.root / "page.docx", rows, **kwargs)

    def run_import(self, rows, **docx_kwargs):
        return wig_import.import_onenote(self.docx(rows, **docx_kwargs), self.assembly_dir, self.workbook)

    def assembly_rows(self):
        ws = load_workbook(self.workbook)["Assembly"]
        return [[cell.value for cell in row] for row in ws.iter_rows(min_row=2)]


THREE = [("Place the base", png_bytes(RED)), ("Fit the bracket", png_bytes(GREEN)), ("Tighten", png_bytes(BLUE))]


class ReadingTests(ImportTestCase):
    def test_steps_numbered_in_row_order_and_padding_skipped(self):
        result = self.run_import(THREE, padding_rows=50)
        self.assertEqual(result.steps, 3)
        self.assertEqual(result.title, "Test page")
        self.assertEqual(self.assembly_rows(),
                         [[1, "Place the base", "Place the base"], [2, "Fit the bracket", "Fit the bracket"],
                          [3, "Tighten", "Tighten"]])

    def test_each_photo_lands_on_its_own_step(self):
        self.run_import(THREE)
        for number, color in enumerate([RED, GREEN, BLUE], start=1):
            self.assertTrue(near(dominant(self.images / f"{number}.jpg"), color), number)

    def test_columns_are_found_by_content_not_position(self):
        self.run_import(THREE, image_first=True)
        self.assertEqual([row[1] for row in self.assembly_rows()], ["Place the base", "Fit the bracket", "Tighten"])
        self.assertTrue(near(dominant(self.images / "2.jpg"), GREEN))

    def test_multi_line_and_nonbreaking_space_text_is_cleaned(self):
        self.run_import([("Tight w   W1/2\n", png_bytes())])
        self.assertEqual(self.assembly_rows()[0][1], "Tight w W1/2")

    def test_photo_without_text_is_a_warning_and_a_blank_step(self):
        result = self.run_import([("First", png_bytes()), (None, png_bytes(GREEN))])
        self.assertEqual(len(result.warnings), 1)
        self.assertIn("row 2", result.warnings[0])
        self.assertIn(self.assembly_rows()[1][1], (None, ""))


class ProblemTests(ImportTestCase):
    def assert_nothing_written(self):
        self.assertFalse(self.workbook.exists())
        self.assertFalse(self.images.exists())

    def test_text_only_row_stops_the_import_and_writes_nothing(self):
        with self.assertRaises(ValueError) as caught:
            self.run_import([("First", png_bytes()), ("A step with no photo", None)])
        self.assertIn("row 2", str(caught.exception))
        self.assertIn("no photo", str(caught.exception))
        self.assert_nothing_written()

    def test_two_photos_in_a_row_stops_the_import(self):
        with self.assertRaises(ValueError) as caught:
            self.run_import([("Two", [png_bytes(), png_bytes()])])
        self.assertIn("2 photos", str(caught.exception))
        self.assert_nothing_written()

    def test_no_table_explains_how_to_fix_it(self):
        with self.assertRaises(ValueError) as caught:
            self.run_import([], with_table=False)
        self.assertIn("Insert > Table", str(caught.exception))

    def test_a_table_with_no_photos_is_not_accepted(self):
        with self.assertRaises(ValueError):
            self.run_import([("only text", None)])

    def test_not_a_docx(self):
        bad = self.root / "notes.docx"
        bad.write_text("plain text", encoding="utf-8")
        with self.assertRaises(ValueError) as caught:
            wig_import.import_onenote(bad, self.assembly_dir, self.workbook)
        self.assertIn("not a readable Word", str(caught.exception))

    def test_unreadable_picture_stops_the_import(self):
        with self.assertRaises(ValueError) as caught:
            self.run_import([("Step", b"not an image")])
        self.assertIn("could not be read", str(caught.exception))
        self.assert_nothing_written()

    def test_bad_options_are_rejected(self):
        docx = self.docx(THREE)
        with self.assertRaises(ValueError):
            wig_import.import_onenote(docx, self.assembly_dir, self.workbook, image_format="gif")
        with self.assertRaises(ValueError):
            wig_import.import_onenote(docx, self.assembly_dir, self.workbook, max_px=10)


class OverwriteTests(ImportTestCase):
    def test_existing_steps_are_not_overwritten_without_replace(self):
        self.run_import(THREE)
        before = self.assembly_rows()
        with self.assertRaises(ValueError) as caught:
            self.run_import([("Different", png_bytes(BLUE))])
        self.assertIn("--replace", str(caught.exception))
        self.assertEqual(self.assembly_rows(), before)

    def test_replace_swaps_steps_and_removes_stale_photos(self):
        self.run_import(THREE)
        (self.images / "1.png").write_bytes(png_bytes(BLUE))  # would shadow 1.jpg: the generator tries PNG first
        result = wig_import.import_onenote(
            self.docx([("Only step", png_bytes(GREEN))]), self.assembly_dir, self.workbook, replace=True)
        self.assertEqual(result.removed_images, 4)
        self.assertEqual(sorted(p.name for p in self.images.iterdir()), ["1.jpg"])
        self.assertEqual(len(self.assembly_rows()), 1)

    def test_existing_workbook_is_backed_up_and_other_sheets_kept(self):
        self.run_import(THREE)
        wb = load_workbook(self.workbook)
        wb["Title Info"]["A2"] = "12-345-67"
        wb.save(self.workbook)
        result = wig_import.import_onenote(
            self.docx([("New", png_bytes())]), self.assembly_dir, self.workbook, replace=True)
        self.assertTrue(result.backup_path.exists())
        self.assertFalse(result.workbook_created)
        self.assertEqual(load_workbook(self.workbook)["Title Info"]["A2"].value, "12-345-67")

    def test_workbook_without_an_assembly_sheet_is_refused(self):
        self.assembly_dir.mkdir(parents=True)
        wb = load_workbook(self._blank_workbook())
        del wb["Assembly"]
        wb.save(self.workbook)
        with self.assertRaises(ValueError) as caught:
            self.run_import(THREE)
        self.assertIn("Assembly", str(caught.exception))

    def _blank_workbook(self) -> Path:
        from wig_scaffold import make_workbook
        path = self.root / "blank.xlsx"
        make_workbook(path)
        return path

    def test_new_workbook_has_all_required_sheets(self):
        result = self.run_import(THREE)
        self.assertTrue(result.workbook_created)
        self.assertEqual(load_workbook(self.workbook).sheetnames,
                         ["Title Info", "Parts", "Equipment", "Assembly", "Testing", "Packaging"])


class SectionTests(ImportTestCase):
    def import_section(self, section, rows, **kwargs):
        return wig_import.import_onenote(
            self.docx(rows), self.assembly_dir, self.workbook, sheet=section, **kwargs)

    def sheet_rows(self, name):
        ws = load_workbook(self.workbook)[name]
        return [[cell.value for cell in row] for row in ws.iter_rows(min_row=2)]

    def test_testing_fills_its_own_sheet_and_folder_only(self):
        self.run_import(THREE)
        result = self.import_section("Testing", [("Check the fit", png_bytes(GREEN)), ("Check torque", png_bytes(RED))])
        self.assertEqual(result.steps, 2)
        self.assertEqual(sorted(p.name for p in (self.assembly_dir / "Testing").iterdir()), ["1.jpg", "2.jpg"])
        self.assertEqual([row[1] for row in self.sheet_rows("Testing")], ["Check the fit", "Check torque"])
        self.assertEqual(len(self.sheet_rows("Assembly")), 3)  # untouched
        self.assertTrue(near(dominant(self.assembly_dir / "Testing" / "1.jpg"), GREEN))

    def test_replace_only_touches_its_own_section(self):
        self.run_import(THREE)
        self.import_section("Packaging", [("Box it", png_bytes(BLUE))])
        self.import_section("Packaging", [("Bag it", png_bytes(GREEN)), ("Label it", png_bytes(RED))], replace=True)
        self.assertEqual([row[1] for row in self.sheet_rows("Packaging")], ["Bag it", "Label it"])
        self.assertEqual(len(list((self.assembly_dir / "Assembly").iterdir())), 3)
        self.assertEqual(len(self.sheet_rows("Assembly")), 3)

    def test_existing_section_is_protected_without_replace(self):
        self.import_section("Testing", [("One", png_bytes())])
        with self.assertRaises(ValueError) as caught:
            self.import_section("Testing", [("Two", png_bytes())])
        self.assertIn("Testing", str(caught.exception))
        self.assertIn("--replace", str(caught.exception))

    def test_a_workbook_without_a_testing_sheet_gets_one(self):
        self.run_import(THREE)
        wb = load_workbook(self.workbook)
        del wb["Testing"]
        wb.save(self.workbook)
        self.import_section("Testing", [("Check", png_bytes())])
        self.assertEqual([row[1] for row in self.sheet_rows("Testing")], ["Check"])

    def test_unknown_section_is_rejected(self):
        with self.assertRaises(ValueError) as caught:
            self.import_section("Inspection", [("x", png_bytes())])
        self.assertIn("Section must be one of", str(caught.exception))


class PhotoTests(unittest.TestCase):
    def test_exif_rotation_is_applied(self):
        # Stored 80x40 with orientation 6 (rotate 90): upright it is 40x80.
        image = wig_import.prepare_image(jpeg_bytes((80, 40), orientation=6), 1600)
        self.assertEqual(image.size, (40, 80))

    def test_large_photos_are_shrunk_and_small_ones_left_alone(self):
        self.assertEqual(max(wig_import.prepare_image(png_bytes(size=(3000, 2000)), 1000).size), 1000)
        self.assertEqual(wig_import.prepare_image(png_bytes(size=(300, 200)), 1000).size, (300, 200))

    def test_transparency_is_flattened_onto_white(self):
        image = wig_import.prepare_image(png_bytes((0, 0, 0, 0), size=(10, 10), mode="RGBA"), 1600)
        self.assertEqual(image.mode, "RGB")
        self.assertEqual(image.getpixel((5, 5)), (255, 255, 255))

    def test_png_option_writes_png(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            docx = make_docx(root / "p.docx", THREE)
            wig_import.import_onenote(docx, root / "a", root / "a" / "a.xlsx", image_format="png")
            self.assertEqual(sorted(p.name for p in (root / "a" / "Assembly").iterdir()), ["1.png", "2.png", "3.png"])


if __name__ == "__main__":
    unittest.main()
