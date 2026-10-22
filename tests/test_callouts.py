"""Tests for safety callouts: CAUTION / WARNING / DANGER / NOTE steps drawn as coloured banners."""

from __future__ import annotations

import copy
import io
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

import generate_work_instruction as wig  # noqa: E402
import wig_glossary  # noqa: E402
import wig_polish  # noqa: E402
from wig_scaffold import SAMPLE_ASSEMBLY, scaffold_project  # noqa: E402


class SplitTests(unittest.TestCase):
    def test_levels_are_recognised_ignoring_case(self):
        cases = {
            "CAUTION: Surface is hot.": ("CAUTION", "Surface is hot."),
            "warning: Pinch point.": ("WARNING", "Pinch point."),
            "Danger:Disconnect power.": ("DANGER", "Disconnect power."),
            "Note: Arrow faces front.": ("NOTE", "Arrow faces front."),
            "CAUTION - Sharp edge.": ("CAUTION", "Sharp edge."),
        }
        for text, expected in cases.items():
            with self.subTest(text):
                self.assertEqual(wig.split_callout(text), expected)

    def test_ordinary_steps_are_not_callouts(self):
        for text in ("Note the orientation of the bracket.", "Cautionary tale.", "Seat the pin. CAUTION: hot.", "", "CAUTION:"):
            with self.subTest(text):
                self.assertEqual(wig.split_callout(text), (None, text))

    def test_polished_notes_become_callouts(self):
        glossary = wig_glossary.build_glossary(copy.deepcopy(wig_glossary.STARTER_GLOSSARY))
        vocab = wig_polish.build_vocabulary(glossary, [])
        draft = wig_polish.polish_note("caution: glue gun nozzle hot NOPART", glossary, [], vocab).draft
        self.assertEqual(wig.split_callout(draft), ("CAUTION", "Glue gun nozzle hot."))


class BuildTests(unittest.TestCase):
    TEXTS = ["DANGER: Disconnect power.", "Place the bracket.", "CAUTION: Glue gun nozzle is hot.", "Note: Arrow faces front."]

    def setUp(self):
        saved = {n: getattr(wig, n) for n in ("DESIGN_FOLDER", "PART_IMAGE_FOLDER", "TEMPLATE_PATH", "FOOTER_ENABLED", "ALLOW_DRAFT")}
        self.addCleanup(lambda: [setattr(wig, n, v) for n, v in saved.items()])
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.made = scaffold_project(Path(tmp.name) / "p", sample=True)
        workbook = self.made["assembly_dir"] / f"{SAMPLE_ASSEMBLY}.xlsx"
        wb = load_workbook(workbook)
        for row, text in enumerate(self.TEXTS, start=2):
            wb["Assembly"].cell(row=row, column=2).value = text
        wb.save(workbook)

    def run_cli(self, *extra):
        argv = ["--assembly", SAMPLE_ASSEMBLY, "--design-folder", str(self.made["design_folder"]),
                "--part-images-folder", str(self.made["part_images_folder"]),
                "--template-path", str(self.made["template_path"]), *extra]
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = wig.main(argv)
        return code, out.getvalue(), err.getvalue()

    def step_slide(self):
        code, _, err = self.run_cli("--initials", "CD", "--welding", "N", "--steps-per-slide", "4", "--initial-release")
        self.assertEqual(code, 0, err)
        deck = self.made["assembly_dir"] / f"{SAMPLE_ASSEMBLY} REV A.pptx"
        return next(s for s in Presentation(deck).slides
                    if any(sh.has_text_frame and "Place the bracket." in sh.text_frame.text for sh in s.shapes))

    def test_callout_steps_become_filled_banners_with_a_bold_label(self):
        slide = self.step_slide()
        callouts = {sh.name: sh for sh in slide.shapes if sh.name.startswith(wig.CALLOUT_NAME)}
        self.assertEqual(sorted(callouts), ["WIG Callout CAUTION", "WIG Callout DANGER", "WIG Callout NOTE"])
        caution = callouts["WIG Callout CAUTION"]
        self.assertEqual(caution.text_frame.text, "CAUTION: Glue gun nozzle is hot.")
        self.assertEqual(caution.fill.fore_color.rgb, wig.CALLOUT_STYLES["CAUTION"][0])
        label, body = caution.text_frame.paragraphs[0].runs
        self.assertEqual((label.text, label.font.bold, body.text), ("CAUTION: ", True, "Glue gun nozzle is hot."))
        self.assertEqual(callouts["WIG Callout NOTE"].text_frame.text, "NOTE: Arrow faces front.")

    def test_ordinary_steps_keep_the_plain_box(self):
        slide = self.step_slide()
        plain = next(sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text == "Place the bracket.")
        self.assertFalse(plain.name.startswith(wig.CALLOUT_NAME))

    def test_check_lists_the_callouts(self):
        code, out, _ = self.run_cli("--check")
        self.assertEqual(code, 0, out)
        self.assertIn("Safety callouts: Assembly 1 (DANGER), Assembly 3 (CAUTION), Assembly 4 (NOTE)", out)


if __name__ == "__main__":
    unittest.main()
