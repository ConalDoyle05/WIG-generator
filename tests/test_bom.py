"""Tests for the BOM cross-check: parts no step names, and part numbers in steps that aren't on the BOM."""

from __future__ import annotations

import io
import json
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

import generate_work_instruction as wig  # noqa: E402
from wig_scaffold import SAMPLE_ASSEMBLY, scaffold_project  # noqa: E402

# Sample Parts sheet: P-001 screw, P-002 washer, P-003 base plate, P-004 bracket, P-005 dowel pin.


class BomCheckTests(unittest.TestCase):
    GLOBALS = ("DESIGN_FOLDER", "PART_IMAGE_FOLDER", "TEMPLATE_PATH", "GLOSSARY_PATH", "FOOTER_ENABLED", "ALLOW_DRAFT")

    def setUp(self):
        saved = {n: getattr(wig, n) for n in self.GLOBALS}
        self.addCleanup(lambda: [setattr(wig, n, v) for n, v in saved.items()])
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.made = scaffold_project(Path(tmp.name) / "p", sample=True)
        self.workbook = self.made["assembly_dir"] / f"{SAMPLE_ASSEMBLY}.xlsx"

    def set_steps(self, descriptions, raw_notes=None):
        wb = load_workbook(self.workbook)
        ws = wb["Assembly"]
        if raw_notes:
            ws["C1"] = "Raw note"
        for row, text in enumerate(descriptions, start=2):
            ws.cell(row=row, column=2).value = text
            if raw_notes:
                ws.cell(row=row, column=3).value = raw_notes[row - 2]
        wb.save(self.workbook)

    def run_cli(self, *extra):
        argv = ["--assembly", SAMPLE_ASSEMBLY, "--design-folder", str(self.made["design_folder"]),
                "--part-images-folder", str(self.made["part_images_folder"]),
                "--template-path", str(self.made["template_path"]), *extra]
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = wig.main(argv)
        return code, out.getvalue(), err.getvalue()

    def bom_lines(self, *extra):
        code, out, err = self.run_cli("--check", *extra)
        self.assertEqual(code, 0, out + err)   # the BOM check never blocks a build
        return [line.strip() for line in out.splitlines() if "BOM check" in line]

    def test_skipped_when_no_step_names_a_part_number(self):
        lines = self.bom_lines()
        self.assertEqual(len(lines), 1)
        self.assertIn("BOM check skipped", lines[0])

    def test_all_parts_used_and_known(self):
        self.set_steps(["Place Base plate (P-003) on the fixture.", "Seat 2 Dowel pin (P-005).",
                        "Set Bracket (P-004) on P-003.", "Fasten with 4 P-001 and 4 P-002."])
        self.assertEqual(self.bom_lines(),
                         ["BOM check: every part on the Parts sheet is named in a step, and every part number "
                          "in the steps is on the Parts sheet."])

    def test_unused_parts_are_listed_with_their_names(self):
        self.set_steps(["Place P-003.", "Seat P-005.", "Set P-004.", "Inspect."])
        self.assertEqual(self.bom_lines(), ["BOM check: 2 part(s) on the Parts sheet are not named in any step: "
                                            "P-001 (M5 x 16 socket head screw), P-002 (M5 washer)."])

    def test_unknown_part_numbers_are_listed_with_where_they_appear(self):
        self.set_steps(["Place P-003 and P-999.", "Seat P-005 and P-999.", "Set P-004 and Q-7.", "Fasten P-001 P-002."])
        lines = self.bom_lines()
        self.assertEqual(lines, ["BOM check: step text names part number(s) that are not on the Parts sheet: "
                                 "P-999 (Assembly 1, Assembly 2); Q-7 (Assembly 3)."])

    def test_matching_ignores_case_and_dash_style_but_not_partial_numbers(self):
        self.set_steps(["Place p-003.", "Seat P\u2013005.", "Set P-004.", "Fasten P-0011 and P-002."])
        lines = self.bom_lines()
        self.assertIn("P-001 (M5 x 16 socket head screw)", lines[0])      # P-0011 is not P-001
        self.assertIn("P-0011", lines[1])
        self.assertNotIn("P-003", " ".join(lines))
        self.assertNotIn("P-005", " ".join(lines))

    def test_the_assembly_and_its_parents_are_not_unknown_parts(self):
        self.set_steps([f"Mark {SAMPLE_ASSEMBLY} on P-003.", "Fit into 98-765-43 with P-005.", "Set P-004.",
                        "Fasten P-001 P-002."])
        self.assertIn("every part on the Parts sheet is named", self.bom_lines()[0])

    def test_a_part_named_only_in_the_floor_note_counts_as_used(self):
        self.set_steps(["Place the base plate.", "Seat P-005.", "Set P-004.", "Fasten P-001 P-002."],
                       raw_notes=["Pl P-003", "Seat P-005", "Set P-004", "Fasten P-001 P-002"])
        self.assertIn("every part on the Parts sheet is named", self.bom_lines()[0])

    def test_the_glossary_part_number_pattern_is_used(self):
        glossary = self.made["glossary"]
        data = json.loads(glossary.read_text(encoding="utf-8"))
        data["settings"]["part_number_pattern"] = r"PN\d{4}"
        glossary.write_text(json.dumps(data), encoding="utf-8")
        self.set_steps(["Place P-003 and PN1234.", "Seat P-005 and Q-7.", "Set P-004.", "Fasten P-001 P-002."])
        lines = self.bom_lines()
        self.assertEqual(lines, ["BOM check: step text names part number(s) that are not on the Parts sheet: "
                                 "PN1234 (Assembly 1)."])

    def test_build_prints_the_check_and_still_succeeds(self):
        self.set_steps(["Place P-003.", "Seat P-005.", "Set P-004.", "Inspect."])
        code, out, err = self.run_cli("--initials", "CD", "--welding", "N", "--steps-per-slide", "4", "--initial-release")
        self.assertEqual(code, 0, err)
        self.assertIn("not named in any step: P-001", out)


if __name__ == "__main__":
    unittest.main()
