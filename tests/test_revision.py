"""Tests for revision tracking: the record saved with each deck, the comparison, and the marks."""

from __future__ import annotations

import io
import json
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
sys.path.insert(0, str(HERE))

import generate_work_instruction as wig  # noqa: E402
import wig_revision  # noqa: E402
from docx_fixture import png_bytes  # noqa: E402
from wig_scaffold import SAMPLE_ASSEMBLY, scaffold_project  # noqa: E402


def record(steps, parts=(), materials=(), tools=(), rev="A"):
    return {
        "format": 1, "revision": rev,
        "sections": {"Assembly": [{"step": str(i), "description": d, "photo": p}
                                  for i, (d, p) in enumerate(steps, start=1)]},
        "parts": [{"number": n, "quantity": q, "name": name} for n, q, name in parts],
        "materials": list(materials), "tools": list(tools),
    }


A, B, C, X = ("Place base.", "a"), ("Seat pin.", "b"), ("Fasten.", "c"), ("Apply glue.", "x")


class CompareTests(unittest.TestCase):
    def test_no_change(self):
        changes = wig_revision.compare(record([A, B, C]), record([A, B, C], rev="B"))
        self.assertFalse(changes.any)
        self.assertEqual(changes.marks, {})
        self.assertEqual(changes.compact_text(), "No content changes")

    def test_inserting_a_step_marks_only_the_new_one(self):
        changes = wig_revision.compare(record([A, B, C]), record([A, X, B, C]))
        self.assertEqual(changes.marks, {("Assembly", "2"): "added"})
        self.assertIn("Assembly 3: same step, renumbered from 2", changes.lines)
        self.assertIn("Assembly 4: same step, renumbered from 3", changes.lines)
        self.assertEqual(changes.compact_text(), "Assembly 2 added")

    def test_removing_a_step(self):
        changes = wig_revision.compare(record([A, B, C]), record([A, C]))
        self.assertEqual(changes.marks, {})
        self.assertIn("Assembly 2 in REV A removed: Seat pin.", changes.lines)
        self.assertEqual(changes.compact_text(), "REV A Assembly 2 removed")

    def test_text_and_photo_changes_are_told_apart(self):
        changes = wig_revision.compare(record([A, B, C]), record([A, ("Seat both pins.", "b"), ("Fasten.", "c2")]))
        self.assertEqual(changes.marks, {("Assembly", "2"): "changed", ("Assembly", "3"): "changed"})
        self.assertEqual(changes.lines, ["Assembly 2: text changed", "Assembly 3: photo changed"])
        self.assertEqual(changes.compact_text(), "Assembly 2, 3 changed")

    def test_swapping_two_steps_is_two_moves_not_two_changes(self):
        changes = wig_revision.compare(record([A, B, C]), record([A, C, B]))
        self.assertEqual(changes.marks, {})
        self.assertEqual(len([l for l in changes.lines if "renumbered" in l]), 2)

    def test_a_new_section_counts_as_added_and_a_dropped_one_as_removed(self):
        old = record([A])
        new = record([A])
        new["sections"]["Testing"] = [{"step": "1", "description": "Check fit.", "photo": "t"}]
        self.assertEqual(wig_revision.compare(old, new).marks, {("Testing", "1"): "added"})
        self.assertIn("Testing 1 in REV A removed: Check fit.", wig_revision.compare(new, old).lines)

    def test_parts_materials_and_tools(self):
        old = record([A], parts=[("P-001", 4, "Screw"), ("P-002", 1, "Washer"), ("P-004", 1, "Bracket")],
                     materials=["Thread locker"], tools=["Torque driver"])
        new = record([A], parts=[("P–001", 6, "Screw"), ("P-003", 1, "Clip"), ("P-004", 1, "L-bracket")],
                     materials=["thread locker"], tools=["Torque driver", "Hex key"])
        changes = wig_revision.compare(old, new)
        self.assertEqual(changes.part_marks, {"P-001", "P-003", "P-004"})
        self.assertIn("Part P–001: quantity 4 → 6", changes.lines)
        self.assertIn("Part P-003 (Clip) added", changes.lines)
        self.assertIn("Part P-002 (Washer) removed", changes.lines)
        self.assertIn("Part P-004: name Bracket → L-bracket", changes.lines)
        self.assertIn("Tools added: Hex key", changes.lines)
        self.assertFalse(any(line.startswith("Materials") for line in changes.lines))   # case-only change

    def test_compact_text_is_kept_short(self):
        steps_old = [(f"Step {i}", str(i)) for i in range(60)]
        steps_new = [(f"Step {i} reworded", str(i)) for i in range(60)]
        text = wig_revision.compare(record(steps_old), record(steps_new)).compact_text()
        self.assertLessEqual(len(text), 150)
        self.assertTrue(text.endswith("…"))

    def test_records_round_trip_and_bad_ones_are_ignored(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "X REV A.wig.json"
            wig_revision.save_manifest(path, record([A]))
            self.assertEqual(wig_revision.load_manifest(path)["sections"]["Assembly"][0]["description"], "Place base.")
            path.write_text("{ not json", encoding="utf-8")
            self.assertIsNone(wig_revision.load_manifest(path))
            self.assertIsNone(wig_revision.load_manifest(Path(tmp) / "missing.wig.json"))
        self.assertEqual(wig_revision.manifest_path(Path("12-345-67 REV A.pptx")).name, "12-345-67 REV A.wig.json")


class BuildTests(unittest.TestCase):
    GLOBALS = ("DESIGN_FOLDER", "PART_IMAGE_FOLDER", "TEMPLATE_PATH", "FOOTER_ENABLED", "ALLOW_DRAFT", "GLOSSARY_PATH")

    def setUp(self):
        saved = {n: getattr(wig, n) for n in self.GLOBALS}
        self.addCleanup(lambda: [setattr(wig, n, v) for n, v in saved.items()])
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.made = scaffold_project(Path(tmp.name) / "p", sample=True)
        self.folder = self.made["assembly_dir"]
        self.workbook = self.folder / f"{SAMPLE_ASSEMBLY}.xlsx"

    def run_cli(self, *extra):
        argv = ["--assembly", SAMPLE_ASSEMBLY, "--design-folder", str(self.made["design_folder"]),
                "--part-images-folder", str(self.made["part_images_folder"]),
                "--template-path", str(self.made["template_path"]), *extra]
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = wig.main(argv)
        return code, out.getvalue(), err.getvalue()

    ANSWERS = ("--initials", "CD", "--welding", "N", "--steps-per-slide", "4")

    def initial(self):
        code, _, err = self.run_cli(*self.ANSWERS, "--initial-release")
        self.assertEqual(code, 0, err)

    def revise(self, rev="B", description="Updated", *extra):
        code, out, err = self.run_cli(*self.ANSWERS, "--revision", rev, "--change-description", description, *extra)
        self.assertEqual(code, 0, err)
        return out

    def edit(self, row, text=None, qty=None, part_row=None):
        wb = load_workbook(self.workbook)
        if text is not None:
            wb["Assembly"].cell(row=row, column=2).value = text
        if qty is not None:
            wb["Parts"].cell(row=part_row, column=3).value = qty
        wb.save(self.workbook)

    def marks(self, rev="B"):
        deck = Presentation(self.folder / f"{SAMPLE_ASSEMBLY} REV {rev}.pptx")
        return [(slide_number, shape.name) for slide_number, slide in enumerate(deck.slides, start=1)
                for shape in slide.shapes if shape.name.startswith(wig.REVISION_MARK_NAME)]

    def test_every_build_saves_a_record_next_to_the_deck(self):
        self.initial()
        saved = json.loads((self.folder / f"{SAMPLE_ASSEMBLY} REV A.wig.json").read_text(encoding="utf-8"))
        self.assertEqual(saved["revision"], "A")
        self.assertEqual(len(saved["sections"]["Assembly"]), 4)
        self.assertEqual(len(saved["parts"]), 5)
        self.assertEqual(len(saved["sections"]["Assembly"][0]["photo"]), 64)

    def test_initial_release_has_no_marks(self):
        self.initial()
        self.assertEqual(self.marks("A"), [])

    def test_a_revision_marks_changed_steps_and_parts_and_archives_the_old_record(self):
        self.initial()
        self.edit(3, text="Set the bracket on the base plate, then check it sits flat.")
        self.edit(None, qty=6, part_row=2)                      # P-001 quantity 4 -> 6
        (self.folder / "Assembly" / "4.png").write_bytes(png_bytes((0, 200, 0)))   # new photo for step 4
        out = self.revise()
        self.assertIn("Changes since REV A (3 marked in the deck):", out)
        self.assertIn("Assembly 2: text changed", out)
        self.assertIn("Assembly 4: photo changed", out)
        self.assertIn("Part P-001: quantity 4 → 6", out)
        names = [name for _, name in self.marks()]
        self.assertEqual(names, ["WIG Revision Mark B"] * 3)
        cancelled = self.folder / "CANCELLED"
        self.assertTrue((cancelled / f"{SAMPLE_ASSEMBLY} REV A.wig.json").exists())
        saved = json.loads((self.folder / f"{SAMPLE_ASSEMBLY} REV B.wig.json").read_text(encoding="utf-8"))
        self.assertEqual(saved["changes_since"], "A")
        self.assertIn("Assembly 2: text changed", saved["changes"])

    def test_marks_are_on_the_right_photos(self):
        self.initial()
        self.edit(3, text="Set the bracket on the base plate, then check it sits flat.")
        self.revise()
        deck = Presentation(self.folder / f"{SAMPLE_ASSEMBLY} REV B.pptx")
        slide = next(s for s in deck.slides if any(sh.name.startswith(wig.REVISION_MARK_NAME) for sh in s.shapes))
        mark = next(sh for sh in slide.shapes if sh.name.startswith(wig.REVISION_MARK_NAME))
        badge = next(sh for sh in slide.shapes if sh.name == f"{wig.STEP_BADGE_NAME} 2")
        self.assertLess(abs(mark.top - badge.top), wig.Inches(0.1))          # same photo, top edge
        self.assertGreater(mark.left, badge.left)                            # right-hand corner

    def test_auto_change_description_goes_in_the_revision_log(self):
        self.initial()
        self.edit(3, text="Set the bracket square to the base plate.")
        self.revise("B", "auto")
        deck = Presentation(self.folder / f"{SAMPLE_ASSEMBLY} REV B.pptx")
        table = wig.find_table(deck.slides[wig.REVISION_LOG_SLIDE_INDEX], "revision log")
        self.assertEqual(table.cell(2, 2).text, "Assembly 2 changed")

    def test_auto_needs_a_previous_record(self):
        self.initial()
        (self.folder / f"{SAMPLE_ASSEMBLY} REV A.wig.json").unlink()
        code, _, err = self.run_cli(*self.ANSWERS, "--revision", "B", "--change-description", "auto")
        self.assertEqual(code, 1)
        self.assertIn("needs the revision record", err)

    def test_a_deck_without_a_record_says_changes_cant_be_marked(self):
        self.initial()
        (self.folder / f"{SAMPLE_ASSEMBLY} REV A.wig.json").unlink()
        self.edit(3, text="Different text.")
        out = self.revise()
        self.assertIn("No revision record for", out)
        self.assertEqual(self.marks(), [])
        self.assertTrue((self.folder / f"{SAMPLE_ASSEMBLY} REV B.wig.json").exists())

    def test_no_changes_is_said_plainly(self):
        self.initial()
        out = self.revise()
        self.assertIn("No changes to steps, parts, materials or tools since REV A.", out)

    def test_a_draft_previews_the_marks_without_saving_a_record(self):
        self.initial()
        self.edit(2, text="Place the base plate on the fixture.")
        code, _, err = self.run_cli(*self.ANSWERS, "--revision", "B", "--change-description", "x", "--draft")
        self.assertEqual(code, 0, err)
        deck = Presentation(self.folder / f"{SAMPLE_ASSEMBLY} REV B DRAFT.pptx")
        self.assertTrue(any(sh.name.startswith(wig.REVISION_MARK_NAME) for s in deck.slides for sh in s.shapes))
        self.assertFalse((self.folder / f"{SAMPLE_ASSEMBLY} REV B DRAFT.wig.json").exists())

    def test_check_previews_what_a_revision_would_mark(self):
        self.initial()
        self.edit(5, text="Inspect for gaps and loose hardware, then wipe it down.")
        code, out, _ = self.run_cli("--check")
        self.assertEqual(code, 0, out)
        self.assertIn("Changes since REV A (a revision would mark them):", out)
        self.assertIn("Assembly 4: text changed", out)

    def test_a_numeric_revision_mark_is_not_mistaken_for_the_page_number(self):
        self.initial()
        self.edit(3, text="Different text.")
        self.revise("2")
        deck = Presentation(self.folder / f"{SAMPLE_ASSEMBLY} REV 2.pptx")
        for number, slide in enumerate(deck.slides, start=1):
            pages = [sh.text_frame.text for sh in slide.shapes if sh.name == wig.PAGE_NUMBER_NAME]
            self.assertEqual(pages, [f"Page {number} of {len(deck.slides)}"])
        self.assertEqual([n for _, n in self.marks("2")], ["WIG Revision Mark 2"])


if __name__ == "__main__":
    unittest.main()
