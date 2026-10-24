"""Tests for the "conventions" section of --config: renamed sheets, moved columns, other markers."""

from __future__ import annotations

import io
import json
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Inches

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
sys.path.insert(0, str(HERE))

import generate_work_instruction as wig  # noqa: E402
from docx_fixture import make_docx, png_bytes  # noqa: E402
from wig_scaffold import SAMPLE_ASSEMBLY, scaffold_project  # noqa: E402

RENAMED = {"title_info": "Cover", "parts": "BOM", "equipment": "Kit", "assembly": "Steps",
           "testing": "Checks", "packaging": "Pack"}
DEFAULT_SHEETS = {"title_info": "Title Info", "parts": "Parts", "equipment": "Equipment", "assembly": "Assembly",
                  "testing": "Testing", "packaging": "Packaging"}


class ConventionTests(unittest.TestCase):
    RUNTIME = ("FOOTER_ENABLED", "ALLOW_DRAFT")

    def setUp(self):
        saved = {n: getattr(wig, n) for n in list(wig.CONFIGURABLE_GLOBALS) + list(self.RUNTIME)}
        self.addCleanup(lambda: [setattr(wig, n, v) for n, v in saved.items()])
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)
        self.made = scaffold_project(self.root / "p", sample=True)
        self.folder = self.made["assembly_dir"]
        self.workbook = self.folder / f"{SAMPLE_ASSEMBLY}.xlsx"

    def config(self, conventions) -> Path:
        path = self.root / "wig.json"
        path.write_text(json.dumps({"conventions": conventions}), encoding="utf-8")
        return path

    def run_cli(self, conventions, *extra):
        argv = ["--assembly", SAMPLE_ASSEMBLY, "--config", str(self.config(conventions)),
                "--design-folder", str(self.made["design_folder"]),
                "--part-images-folder", str(self.made["part_images_folder"]),
                "--template-path", str(self.made["template_path"]), *extra]
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = wig.main(argv)
        return code, out.getvalue(), err.getvalue()

    def build(self, conventions, *extra):
        return self.run_cli(conventions, "--initials", "CD", "--welding", "N", "--steps-per-slide", "4",
                            "--initial-release", *extra)

    def texts(self, slide_index=None):
        deck = Presentation(self.folder / f"{SAMPLE_ASSEMBLY} REV A.pptx")
        slides = deck.slides if slide_index is None else [deck.slides[slide_index]]
        return [sh.text_frame.text for s in slides for sh in s.shapes if sh.has_text_frame]

    def rename_sheets(self):
        wb = load_workbook(self.workbook)
        for role, name in RENAMED.items():
            wb[DEFAULT_SHEETS[role]].title = name
        wb.save(self.workbook)

    def test_renamed_sheets_build(self):
        self.rename_sheets()
        code, _, err = self.build({"sheets": RENAMED})
        self.assertEqual(code, 0, err)
        texts = self.texts()
        self.assertIn("Assembly Steps 1\u20134", texts)
        self.assertTrue(any("QTY: 4" in t for t in texts))

    def test_renamed_sheets_without_the_config_say_which_sheet_is_missing(self):
        self.rename_sheets()
        code, _, err = self.build({})
        self.assertEqual(code, 1)
        self.assertIn("Title Info", err)

    def test_moved_columns(self):
        wb = load_workbook(self.workbook)
        ws = wb["Parts"]
        rows = [[c.value for c in r] for r in ws.iter_rows(min_row=1, max_col=4)]
        ws.delete_rows(1, ws.max_row)
        for number, bom, qty, name in rows:          # Part Name | Quantity | BOMNAME | Part #
            ws.append([name, qty, bom, number])
        title = wb["Title Info"]
        values = [c.value for c in title[2]]
        title.delete_rows(2)
        title.append([None, *values])                 # everything one column to the right
        wb.save(self.workbook)
        conventions = {
            "parts_columns": {"part_name": 1, "quantity": 2, "bomname": 3, "part_number": 4},
            "title_info_columns": {"WI #": 2, "WI Title": 3, "REV #": 4, "PRODUCTION CELL": 5, "FORMER WI#": 6,
                                   "ASSEMBLY TIME": 7, "Parent Assemblies": 8},
        }
        code, _, err = self.build(conventions)
        self.assertEqual(code, 0, err)
        texts = self.texts()
        self.assertIn("Demo Bracket Assembly", texts)
        self.assertTrue(any(t.startswith("P-001  |  QTY: 4") for t in texts))

    def test_optional_title_fields(self):
        wb = load_workbook(self.workbook)
        wb["Title Info"]["D2"] = None      # PRODUCTION CELL
        wb["Title Info"]["F2"] = None      # ASSEMBLY TIME, with no step times
        wb.save(self.workbook)
        code, _, err = self.build({})
        self.assertEqual(code, 1)
        self.assertIn("PRODUCTION CELL", err)
        code, _, err = self.build({"required_title_fields": ["WI #", "WI Title", "REV #"]})
        self.assertEqual(code, 0, err)

    def test_custom_subassembly_marker(self):
        wb = load_workbook(self.workbook)
        for row in wb["Parts"].iter_rows(min_row=2):
            if row[1].value == "MFG BOM":
                row[1].value = "Subassy"
        wb.save(self.workbook)
        code, _, err = self.build({"subassembly_marker": "subassy"})
        self.assertEqual(code, 0, err)
        self.assertIn("Includes: P-003, P-004", self.texts(0))

    def test_renamed_folders(self):
        shutil.move(str(self.folder / "Assembly"), str(self.folder / "Assembly photos"))
        conventions = {"folders": {"assembly": "Assembly photos", "cancelled": "Superseded"}}
        code, _, err = self.build(conventions)
        self.assertEqual(code, 0, err)
        code, _, err = self.run_cli(conventions, "--initials", "CD", "--welding", "N", "--steps-per-slide", "4",
                                    "--revision", "B", "--change-description", "x")
        self.assertEqual(code, 0, err)
        self.assertTrue((self.folder / "Superseded" / f"{SAMPLE_ASSEMBLY} REV A.pptx").exists())
        self.assertFalse((self.folder / "CANCELLED").exists())

    def test_custom_placeholders_and_markers(self):
        prs = Presentation(self.made["template_path"])
        cover = prs.slides[0]
        cover.shapes.add_textbox(Inches(0.5), Inches(6.2), Inches(5), Inches(0.5)).text_frame.text = "Name: <<TITLE>>"
        prs.slides[2].shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text_frame.text = "Torque Chart\nN/A IF EMPTY"
        prs.save(self.made["template_path"])
        conventions = {"placeholders": {"title": "<<TITLE>>"}, "blank_page_marker": "n/a if empty",
                       "optional_page_titles": ["torque chart"]}
        code, _, err = self.build(conventions)
        self.assertEqual(code, 0, err)
        texts = self.texts()
        self.assertIn("Name: Demo Bracket Assembly", texts)
        self.assertFalse(any("Torque Chart" in t for t in texts))

    def test_import_and_polish_follow_renamed_sheets_and_columns(self):
        self.rename_sheets()
        wb = load_workbook(self.workbook)
        ws = wb["BOM"]
        rows = [[c.value for c in r] for r in ws.iter_rows(min_row=1, max_col=4)]
        ws.delete_rows(1, ws.max_row)
        for number, bom, qty, name in rows:
            ws.append([name, qty, bom, number])
        wb.save(self.workbook)
        conventions = {"sheets": RENAMED, "parts_columns": {"part_name": 1, "quantity": 2, "bomname": 3, "part_number": 4},
                       "folders": {"assembly": "Step photos"}}
        page = make_docx(self.root / "page.docx", [("Tight w W1/2 P-001", png_bytes()), ("Seat PN?", png_bytes())])
        code, out, err = self.run_cli(conventions, "--import-onenote", str(page), "--replace", "--polish")
        self.assertEqual(code, 0, err)
        self.assertTrue((self.folder / "Step photos" / "1.jpg").exists())
        steps = load_workbook(self.workbook)["Steps"]
        self.assertEqual(steps.cell(row=2, column=2).value,
                         'Tighten with a 1/2" wrench M5 x 16 socket head screw (P-001).')
        self.assertEqual(steps.cell(row=3, column=4).value, "Flagged")

    def test_bad_conventions_are_rejected_and_change_nothing(self):
        before = {n: getattr(wig, n) for n in wig.CONFIGURABLE_GLOBALS}
        cases = {
            "unknown key": ({"sheet_names": {}}, 'Unknown config "conventions" key'),
            "unknown sheet role": ({"sheets": {"bom": "BOM"}}, "Unknown conventions sheets key"),
            "blank sheet": ({"sheets": {"parts": " "}}, "non-empty text"),
            "excel-illegal sheet": ({"sheets": {"parts": "Parts/BOM"}}, "isn't allowed by Excel"),
            "long sheet": ({"sheets": {"parts": "P" * 32}}, "isn't allowed by Excel"),
            "duplicate sheets": ({"sheets": {"parts": "Assembly"}}, "share one sheet name"),
            "column zero": ({"parts_columns": {"part_number": 0}}, "column number from 1 to 50"),
            "shared column": ({"parts_columns": {"part_number": 2}}, "share one column"),
            "unknown title field": ({"required_title_fields": ["WI #", "REV #", "Color"]}, "Unknown title field"),
            "REV # dropped": ({"required_title_fields": ["WI #"]}, '"REV #" must stay required'),
            "folder path": ({"folders": {"cancelled": "old/CANCELLED"}}, "single folder name"),
            "shared folder": ({"folders": {"testing": "Assembly"}}, "share one folder name"),
            "shared placeholder": ({"placeholders": {"title": "{REV}"}}, "share one placeholder"),
            "empty marker list": ({"step_markers": []}, "must be a list of text"),
            "marker not text": ({"subassembly_marker": 5}, "non-empty text"),
        }
        for label, (conventions, message) in cases.items():
            with self.subTest(label):
                with self.assertRaises(ValueError) as caught:
                    wig.load_config(self.config(conventions))
                self.assertIn(message, str(caught.exception))
                self.assertEqual({n: getattr(wig, n) for n in wig.CONFIGURABLE_GLOBALS}, before)

    def test_slides_and_conventions_combine(self):
        wig.load_config(self._both())
        assembly = wig.PROCESS_SECTIONS[0]
        self.assertEqual((assembly.sheet, assembly.folder, assembly.template_index), ("Steps", "Step photos", 7))

    def _both(self) -> Path:
        path = self.root / "both.json"
        path.write_text(json.dumps({"slides": {"assembly": 7, "testing": 8, "packaging": 9},
                                    "conventions": {"sheets": {"assembly": "Steps"},
                                                    "folders": {"assembly": "Step photos"}}}), encoding="utf-8")
        return path


if __name__ == "__main__":
    unittest.main()
