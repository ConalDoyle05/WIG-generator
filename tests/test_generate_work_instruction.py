"""Tests for generate_work_instruction.py.

Loaded by path (not import) the same way the script is meant to be run
standalone, so the test suite doesn't require the project to be packaged.
"""

from __future__ import annotations

import importlib.util
import io
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

MODULE_PATH = Path(__file__).parents[1] / "generate_work_instruction.py"
SPEC = importlib.util.spec_from_file_location("generate_work_instruction", MODULE_PATH)
assert SPEC and SPEC.loader
wig = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = wig
SPEC.loader.exec_module(wig)


def make_png(path: Path, size=(20, 20), color=(200, 30, 30)) -> None:
    Image.new("RGB", size, color).save(path)


# =============================================================================
# Pure-logic unit tests
# =============================================================================

class CleanCellTests(unittest.TestCase):
    def test_none_becomes_empty_string(self):
        self.assertEqual(wig.clean_cell(None), "")

    def test_trailing_float_zero_stripped(self):
        self.assertEqual(wig.clean_cell(12.0), "12")

    def test_non_integer_float_kept_as_is(self):
        self.assertEqual(wig.clean_cell(12.5), "12.5")

    def test_datetime_formatted_mmddyyyy(self):
        self.assertEqual(wig.clean_cell(datetime(2026, 3, 4)), "03/04/2026")


class NormalizeTests(unittest.TestCase):
    def test_normalize_number_uppercases_and_strips_spaces(self):
        self.assertEqual(wig.normalize_number(" 12-345-67 "), "12-345-67")

    def test_normalize_header_strips_punctuation(self):
        self.assertEqual(wig.normalize_header("Part #"), "PART")


class FormatListTests(unittest.TestCase):
    def test_format_parent_list_empty_is_none(self):
        self.assertEqual(wig.format_parent_list(""), "None")

    def test_format_parent_list_splits_on_comma_and_semicolon(self):
        self.assertEqual(wig.format_parent_list("A, B; C"), "A, B, C")

    def test_format_included_subassemblies_empty_is_none(self):
        data = wig.WorkInstructionData(
            assembly_number="X", reviser_initials="CD", uses_welding=False,
            change_description="", is_revision=False, new_revision=None,
            title_info={}, parts=[], included_subassemblies=[], materials=[],
            tools=[], process_steps={},
        )
        self.assertEqual(wig.format_included_subassemblies(data), "None")


class RevisionSortTests(unittest.TestCase):
    def test_split_revision_with_dot(self):
        self.assertEqual(wig.split_revision("D.2"), ("D", "2"))

    def test_split_revision_without_dot(self):
        self.assertEqual(wig.split_revision("D"), ("D", ""))

    def test_revision_sort_key_orders_numeric_revision_not_lexically(self):
        # REV D.10 must sort after REV D.9 (numeric), not before it (lexical).
        p9 = Path("12-345-67 REV D.9.pptx")
        p10 = Path("12-345-67 REV D.10.pptx")
        self.assertLess(wig.revision_sort_key(p9), wig.revision_sort_key(p10))


class ImagePathTests(unittest.TestCase):
    def test_exact_match_found(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            make_png(folder / "ABC123.png")
            found = wig.image_path_from_folder(folder, "abc123")
            self.assertEqual(found.name, "ABC123.png")

    def test_case_insensitive_fallback(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            make_png(folder / "Part.JPG")
            found = wig.image_path_from_folder(folder, "part")
            # Windows and default macOS volumes ignore case, so the exact-name
            # check already finds the file under the requested casing there.
            self.assertTrue(found.exists())
            self.assertEqual(found.name.lower(), "part.jpg")

    def test_missing_image_returns_expected_default_path(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            found = wig.image_path_from_folder(folder, "missing")
            self.assertEqual(found, folder / "MISSING.png")


class UniqueDestinationTests(unittest.TestCase):
    def test_no_collision_returns_same_name(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            dest = wig.unique_destination(folder, "file.pptx")
            self.assertEqual(dest.name, "file.pptx")

    def test_collision_adds_numeric_suffix(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            (folder / "file.pptx").touch()
            dest = wig.unique_destination(folder, "file.pptx")
            self.assertEqual(dest.name, "file (2).pptx")


class DescriptionBoxGrowthTests(unittest.TestCase):
    def test_short_description_keeps_standard_box(self):
        box = (0.45, 1.05, 4.45, 0.85)
        result = wig.adjust_description_box(box, "short text", bottom_limit=7.30)
        self.assertEqual(result, box)

    def test_long_description_grows_downward_only(self):
        box = (0.45, 1.05, 4.45, 0.85)
        long_text = "x" * 200
        x, y, w, h = wig.adjust_description_box(box, long_text, bottom_limit=7.30)
        self.assertEqual((x, y, w), (0.45, 1.05, 4.45))
        self.assertGreater(h, 0.85)

    def test_long_description_stops_at_bottom_limit(self):
        box = (0.45, 6.90, 4.45, 0.30)
        long_text = "x" * 200
        _, y, _, h = wig.adjust_description_box(box, long_text, bottom_limit=7.30)
        self.assertLessEqual(y + h, 7.30 + 1e-9)

    def test_description_bottom_limit_stops_above_the_box_below_it(self):
        layout_boxes = [
            ((0, 0, 1, 1), (0, 1.0, 1, 0.5)),
            ((0, 2.0, 1, 1), (0, 3.0, 1, 0.5)),
        ]
        limit = wig.description_bottom_limit(layout_boxes, 0)
        self.assertEqual(limit, 2.0 - wig.LONG_DESCRIPTION_ROW_GAP)


class ResolveColumnsTests(unittest.TestCase):
    def test_warns_when_alias_header_found_in_wrong_column(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Quantity", "Part #", "BOMNAME", "Part Name"])
        buf = io.StringIO()
        with redirect_stdout(buf):
            wig.resolve_columns(ws, "Parts", wig.PARTS_COLUMNS)
        self.assertIn("Warning", buf.getvalue())

    def test_no_warning_for_correctly_ordered_headers(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Part #", "BOMNAME", "Quantity", "Part Name"])
        buf = io.StringIO()
        with redirect_stdout(buf):
            wig.resolve_columns(ws, "Parts", wig.PARTS_COLUMNS)
        self.assertEqual(buf.getvalue(), "")

    def test_raises_when_sheet_has_too_few_columns(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Only", "Two"])
        with self.assertRaises(ValueError):
            wig.resolve_columns(ws, "Parts", wig.PARTS_COLUMNS)


# =============================================================================
# Integration tests: build a real deck from a synthetic template
# =============================================================================

def build_synthetic_template(path: Path) -> None:
    """Builds a minimal template matching the fixed slide-index contract.

    Slide 4 (the parts template) deliberately carries a stray picture, which
    is exactly the setup that exposed the duplicate-slide picture bug: any
    picture already on a template slide must survive duplication cleanly.
    """
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]
    add_slide = lambda: prs.slides.add_slide(blank_layout)

    # 0: title
    slide0 = add_slide()
    slide0.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)) \
        .text_frame.text = "{TITLE} - {Assembly Number} REV {REV}"
    slide0.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1)) \
        .text_frame.text = "Parents: {PARENT_ASSEMBLIES}"

    # 1: revision log
    slide1 = add_slide()
    table = slide1.shapes.add_table(6, 5, Inches(1), Inches(1), Inches(8), Inches(4)).table
    for col, header in enumerate(["DWG REV", "WI REV", "Change", "By", "Date"]):
        table.cell(0, col).text = header

    # 2: filler (not referenced by any fixed index)
    add_slide()

    # 3: welding
    add_slide()

    # 4: parts template, with a stray picture to regression-test duplication
    slide4 = add_slide()
    stray_image = path.parent / "_stray.png"
    make_png(stray_image, size=(10, 10), color=(0, 0, 0))
    slide4.shapes.add_picture(str(stray_image), Inches(0.1), Inches(0.1), Inches(0.3), Inches(0.3))
    slide4.shapes.add_textbox(Inches(0.5), Inches(4), Inches(2), Inches(0.5)) \
        .text_frame.text = "{Part Number}"

    # 5: equipment
    slide5 = add_slide()
    slide5.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(3), Inches(0.5)) \
        .text_frame.text = "{Materials placeholder} {Tools placeholder}"

    # 6: assembly steps
    slide6 = add_slide()
    slide6.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(3), Inches(0.5)) \
        .text_frame.text = "{Assembly Step picture} {Step Description}"

    # 7: filler
    add_slide()
    # 8: testing (optional)
    add_slide()
    # 9: packaging (optional)
    add_slide()

    prs.save(path)


def find_shape_with_text(prs: Presentation, needle: str):
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and needle in shape.text_frame.text:
                return shape
    return None


def slide_has_text(slide, needle: str) -> bool:
    return any(
        shape.has_text_frame and needle in shape.text_frame.text
        for shape in slide.shapes
    )


def count_pictures(slide) -> int:
    return sum(1 for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)


class ProcessSlideTitleTests(unittest.TestCase):
    spec = wig.PROCESS_SECTIONS[0]

    def rows(self, *numbers):
        return [wig.ProcessRow(step_number=n, description="x") for n in numbers]

    def test_range_for_several_steps(self):
        self.assertEqual(wig.process_slide_title(self.spec, self.rows("1", "2", "3", "4")),
                         "Assembly Steps 1\u20134")

    def test_single_step(self):
        self.assertEqual(wig.process_slide_title(self.spec, self.rows("9")), "Assembly Step 9")

    def test_testing_and_packaging_are_not_pluralised_badly(self):
        testing, packaging = wig.PROCESS_SECTIONS[1], wig.PROCESS_SECTIONS[2]
        self.assertEqual(wig.process_slide_title(testing, self.rows("1", "2")), "Testing Steps 1\u20132")
        self.assertEqual(wig.process_slide_title(testing, self.rows("3")), "Testing Step 3")
        self.assertEqual(wig.process_slide_title(packaging, self.rows("1", "2", "3")), "Packaging Steps 1\u20133")

    def test_range_uses_the_workbook_step_numbers(self):
        self.assertEqual(wig.process_slide_title(self.spec, self.rows("5", "7")), "Assembly Steps 5\u20137")


class BuildWorkInstructionIntegrationTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        root = Path(self._tmp.name)

        self.assembly_number = "12-345-67"
        self.part_images = root / "Part_Catalog_Images"
        self.part_images.mkdir()
        self.design_folder = root / "Design Folder"
        self.assembly_dir = self.design_folder / "12" / self.assembly_number
        (self.assembly_dir / "Assembly").mkdir(parents=True)

        self.part_numbers = [f"P-{index:03d}" for index in range(12)]
        for part_number in self.part_numbers:
            make_png(self.part_images / f"{part_number}.png")
        make_png(self.part_images / f"{self.assembly_number}.png")
        make_png(self.assembly_dir / "Assembly" / "1.png")
        make_png(self.assembly_dir / "Assembly" / "2.png")

        self.template_path = root / "template.pptx"
        build_synthetic_template(self.template_path)

        self._original_part_folder = wig.PART_IMAGE_FOLDER
        self._original_design_folder = wig.DESIGN_FOLDER
        self._original_template_path = wig.TEMPLATE_PATH
        wig.PART_IMAGE_FOLDER = self.part_images
        wig.DESIGN_FOLDER = self.design_folder
        wig.TEMPLATE_PATH = self.template_path

    def tearDown(self):
        wig.PART_IMAGE_FOLDER = self._original_part_folder
        wig.DESIGN_FOLDER = self._original_design_folder
        wig.TEMPLATE_PATH = self._original_template_path

    def _data(self, **overrides) -> "wig.WorkInstructionData":
        parts = [
            wig.PartRow(part_number=pn, bomname="PURCHASED", quantity="1", part_name=f"Widget {pn}")
            for pn in self.part_numbers
        ]
        steps = {
            wig.SHEET_ASSEMBLY_STEPS: [
                wig.ProcessRow(step_number="1", description="First step"),
                wig.ProcessRow(step_number="2", description="Second step"),
            ],
            wig.SHEET_TESTING: [],
            wig.SHEET_PACKAGING: [],
        }
        defaults = dict(
            assembly_number=self.assembly_number,
            reviser_initials="CD",
            uses_welding=False,
            change_description="Initial release",
            is_revision=False,
            new_revision=None,
            title_info={
                "WI #": self.assembly_number,
                "WI Title": "Test Widget",
                "REV #": "A",
                "PRODUCTION CELL": "Cell 1",
                "FORMER WI#": "",
                "ASSEMBLY TIME": "10 min",
                "DATE": "01/01/2026",
                "Parent Assemblies": "",
            },
            parts=parts,
            included_subassemblies=[],
            materials=["Grease"],
            tools=["Torque wrench"],
            process_steps=steps,
        )
        defaults.update(overrides)
        return wig.WorkInstructionData(**defaults)

    def test_initial_release_generates_expected_output(self):
        data = self._data()
        output_path, archive = wig.build_work_instruction(data, items_per_slide=2)

        self.assertTrue(output_path.exists())
        self.assertEqual(archive.moved, [])

        prs = Presentation(output_path)

        # Regression: a duplicated template slide's own picture must not
        # survive as a broken/orphaned image, and each parts slide should
        # carry exactly its own part-card pictures (10 on the first, 2 on
        # the second — 12 parts split across PARTS_PER_SLIDE=10).
        parts_slides = [slide for slide in prs.slides if slide_has_text(slide, "QTY:")]
        self.assertEqual(len(parts_slides), 2)
        self.assertEqual(sorted(count_pictures(slide) for slide in parts_slides), [2, 10])

        # Regression: every run in multi-line text (not just the first line)
        # must carry the requested font size.
        card_shape = find_shape_with_text(prs, "QTY: 1")
        self.assertIsNotNone(card_shape)
        card_runs = card_shape.text_frame.paragraphs[0].runs
        self.assertGreaterEqual(len(card_runs), 2)
        for run in card_runs:
            self.assertEqual(run.font.size, Pt(13))

        materials_shape = find_shape_with_text(prs, "- Grease")
        self.assertIsNotNone(materials_shape)
        materials_runs = materials_shape.text_frame.paragraphs[0].runs
        self.assertGreaterEqual(len(materials_runs), 2)
        for run in materials_runs:
            self.assertEqual(run.font.size, Pt(16))

        # Assembly steps populated; optional sections with no data omitted.
        self.assertIsNotNone(find_shape_with_text(prs, "First step"))
        self.assertIsNone(find_shape_with_text(prs, "Testing #1"))
        self.assertIsNone(find_shape_with_text(prs, "Packaging #1"))

    def test_each_step_picture_carries_its_own_number(self):
        # Regression: renumber_all_slides took the first digit-only shape for the page
        # number, so the "1" badge was moved to the footer and renumbered.
        output_path, _ = wig.build_work_instruction(self._data(), items_per_slide=2)
        prs = Presentation(output_path)
        slide = next(s for s in prs.slides if slide_has_text(s, "First step"))
        badges = sorted((sh for sh in slide.shapes if sh.name.startswith(wig.STEP_BADGE_NAME)),
                        key=lambda sh: sh.left)
        self.assertEqual([b.text_frame.text for b in badges], ["1", "2"])
        pictures = sorted((sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE),
                          key=lambda sh: sh.left)
        for badge, picture in zip(badges, pictures):
            # Sits on the picture's own top-left corner, not the (larger) box corner.
            self.assertEqual(badge.left, picture.left + Inches(wig.STEP_BADGE_INSET))
            self.assertEqual(badge.top, picture.top + Inches(wig.STEP_BADGE_INSET))
        self.assertTrue(slide_has_text(slide, "Assembly Steps 1\u20132"))
        # The page number still exists as its own shape.
        self.assertTrue(any(sh.has_text_frame and not sh.name.startswith(wig.STEP_BADGE_NAME)
                            and sh.text_frame.text.strip().isdigit() for sh in slide.shapes))

    def _template_with_blank_cycle_time_page(self):
        """Puts a blank optional page AHEAD of the welding slide, as a real template can."""
        prs = Presentation(self.template_path)
        prs.slides[2].shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)) \
            .text_frame.text = "Cycle Time\nLeave page blank if no data"
        prs.slides[3].shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)) \
            .text_frame.text = "WELDING-SLIDE"
        prs.save(self.template_path)

    def test_welding_removed_without_deleting_the_parts_list(self):
        # Regression: the blank Cycle Time page (index 2) is deleted first, which
        # shifts the welding slide up. Deleting "slide 3" afterwards removed the
        # parts list instead and left the welding page in the deck.
        self._template_with_blank_cycle_time_page()
        output_path, _ = wig.build_work_instruction(self._data(uses_welding=False), items_per_slide=2)
        prs = Presentation(output_path)
        self.assertIsNone(find_shape_with_text(prs, "WELDING-SLIDE"))
        self.assertIsNone(find_shape_with_text(prs, "Cycle Time"))
        self.assertEqual(len([s for s in prs.slides if slide_has_text(s, "QTY:")]), 2)

    def test_welding_kept_when_used(self):
        self._template_with_blank_cycle_time_page()
        output_path, _ = wig.build_work_instruction(self._data(uses_welding=True), items_per_slide=2)
        prs = Presentation(output_path)
        self.assertIsNotNone(find_shape_with_text(prs, "WELDING-SLIDE"))
        self.assertEqual(len([s for s in prs.slides if slide_has_text(s, "QTY:")]), 2)

    def test_short_template_gives_a_clear_error(self):
        prs = Presentation(self.template_path)
        wig.delete_slide(prs, len(prs.slides) - 1)
        prs.save(self.template_path)
        with self.assertRaises(ValueError) as caught:
            wig.build_work_instruction(self._data(), items_per_slide=2)
        self.assertIn("slides", str(caught.exception))

    def test_save_refuses_to_overwrite_existing_output(self):
        data = self._data()
        wig.build_work_instruction(data, items_per_slide=2)
        with self.assertRaises(FileExistsError):
            wig.build_work_instruction(data, items_per_slide=2)

    def test_revision_preserves_prior_log_row_and_archives_old_deck(self):
        initial = self._data()
        output_path, _ = wig.build_work_instruction(initial, items_per_slide=2)

        revision = self._data(is_revision=True, new_revision="B", change_description="Fixed torque spec")
        revision.title_info["REV #"] = "B"
        output_path_b, archive = wig.build_work_instruction(revision, items_per_slide=2)

        self.assertNotEqual(output_path, output_path_b)
        self.assertEqual(len(archive.moved), 1)
        self.assertEqual(archive.moved[0].parent.name, wig.CANCELLED_FOLDER_NAME)
        self.assertFalse(output_path.exists())
        self.assertTrue(output_path_b.exists())

        prs = Presentation(output_path_b)
        table = wig.find_table(prs.slides[wig.REVISION_LOG_SLIDE_INDEX], "revision log")
        row_texts = [table.cell(row, 0).text for row in range(1, 3)]
        self.assertEqual(row_texts, ["A", "B"])


if __name__ == "__main__":
    unittest.main()
