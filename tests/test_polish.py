"""Tests for the glossary and the polish step."""

from __future__ import annotations

import copy
import csv
import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

import wig_glossary  # noqa: E402
import wig_polish  # noqa: E402
from wig_glossary import STARTER_GLOSSARY, build_glossary  # noqa: E402
from wig_polish import Part, build_vocabulary, polish_note  # noqa: E402
from wig_scaffold import make_workbook  # noqa: E402

PARTS = [Part("P-004", "Bracket"), Part("12-345-67", "Base plate"), Part("P-001", "Dowel pin")]


def glossary(**changes):
    data = copy.deepcopy(STARTER_GLOSSARY)
    for key, value in changes.items():
        if isinstance(value, dict) and isinstance(data.get(key), dict):
            data[key].update(value)
        else:
            data[key] = value
    return build_glossary(data)


class EngineTestCase(unittest.TestCase):
    def setUp(self):
        self.g = glossary()
        self.vocab = build_vocabulary(self.g, PARTS)

    def polish(self, note, g=None, parts=PARTS):
        g = g or self.g
        return polish_note(note, g, parts, build_vocabulary(g, parts) if g is not self.g else self.vocab)

    def assert_draft(self, note, expected, flags=0):
        result = self.polish(note)
        self.assertEqual(result.draft, expected)
        self.assertEqual(len(result.flags), flags, result.flags)
        return result


class ExpansionTests(EngineTestCase):
    def test_terms_tools_and_part_numbers(self):
        result = self.assert_draft("Tight w W1/2 P-004", 'Tighten with a 1/2" wrench Bracket (P-004).')
        self.assertEqual(result.status, "Ready")
        self.assertEqual(result.tools_used, ["W1/2"])
        self.assertEqual(result.parts_used, ["P-004"])

    def test_capitalises_and_ends_the_sentence(self):
        self.assert_draft("insp P-004 for gaps", "Inspect Bracket (P-004) for gaps.")
        self.assert_draft("Insp P-004 for gaps!", "Inspect Bracket (P-004) for gaps!")

    def test_punctuation_stays_attached(self):
        self.assert_draft("Place 12-345-67 on fixture, align w P-001.",
                          "Place Base plate (12-345-67) on fixture, align with Dowel pin (P-001).")

    def test_part_numbers_match_ignoring_case_and_dash_style(self):
        self.assert_draft("Seat p-004 now", "Seat Bracket (P-004) now.")
        self.assert_draft("Seat 12–345–67", "Seat Base plate (12-345-67).")

    def test_part_format_setting(self):
        g = glossary(settings={"part_format": "{number} {name}"})
        self.assertEqual(self.polish("Seat P-004", g).draft, "Seat P-004 Bracket.")

    def test_longest_term_wins(self):
        g = glossary(terms={"tight w": "tighten using"})
        self.assertEqual(self.polish("tight w W1/2 P-004", g).draft, 'Tighten using a 1/2" wrench Bracket (P-004).')

    def test_multi_word_term_does_not_cross_punctuation(self):
        g = glossary(terms={"tight w": "tighten using"})
        self.assertEqual(self.polish("tight, w W1/2 P-004", g).draft, 'Tighten, with a 1/2" wrench Bracket (P-004).')

    def test_tool_codes_ignore_case(self):
        self.assert_draft("tight w w1/2 P-004", 'Tighten with a 1/2" wrench Bracket (P-004).')

    def test_smart_quotes_and_spaces_are_normalised(self):
        self.assert_draft("Slide 1/2”  past edge NOPART", 'Slide 1/2" past edge.')

    def test_already_professional_text_passes_through(self):
        self.assert_draft("Place the Bracket (P-004) on the fixture.", "Place the Bracket (P-004) on the fixture.")

    def test_polishing_a_draft_again_changes_nothing(self):
        for note in ("Tight w W1/2 P-004", "Torque P-004 to 25 Nm", "Rotate 90deg cw NOPART",
                     "Place 12-345-67 on fixture, align w P-001."):
            with self.subTest(note):
                draft = self.polish(note).draft
                again = self.polish(draft + (" NOPART" if "NOPART" in note else "")).draft
                self.assertEqual(again, draft)


class NumberTests(EngineTestCase):
    def test_unit_after_number_is_normalised(self):
        result = self.polish("Torque P-004 to 25 Nm")
        self.assertEqual(result.draft, "Torque Bracket (P-004) to 25 N·m.")
        self.assertFalse(any("unit" in f for f in result.flags))

    def test_unit_attached_to_number(self):
        self.assertEqual(self.polish("Rotate 90deg cw NOPART").draft, "Rotate 90° clockwise.")
        self.assertEqual(self.polish("Torque to 12.5in-lb P-004").draft, "Torque to 12.5 in-lb Bracket (P-004).")

    def test_number_without_unit_is_flagged(self):
        result = self.polish("Torque P-004 to 25")
        self.assertTrue(any(f.startswith("Number without a unit: 25") for f in result.flags), result.flags)
        self.assertEqual(result.status, "Flagged")

    def test_counts_are_fine(self):
        for note in ("Insert 4 screws NOPART", "Insert 2 P-001", "Apply 2 drops NOPART", "Insert 4x screws NOPART",
                     "Insert 2 dowel pins NOPART"):
            with self.subTest(note):
                result = self.polish(note)
                self.assertFalse(any("unit" in f for f in result.flags), result.flags)

    def test_numbers_always_come_through(self):
        result = self.polish("Place 3 brkts 25 mm apart NOPART")
        self.assertIn("3", result.draft)
        self.assertIn("25 mm", result.draft)
        self.assertFalse(any(f.startswith("Internal check") for f in result.flags))

    def test_a_term_that_swallows_a_number_trips_the_guard(self):
        g = glossary(terms={"2x": "twice"})
        result = self.polish("Wipe 2x NOPART", g)
        self.assertTrue(any(f.startswith("Internal check") for f in result.flags), result.flags)


class PartStateTests(EngineTestCase):
    def test_missing_part_number_is_flagged_by_default(self):
        result = self.polish("Tight w W1/2")
        self.assertIn("No part number", " ".join(result.flags))

    def test_pn_placeholder(self):
        result = self.polish("Seat PN? on fixture")
        self.assertEqual(result.draft, "Seat [part number needed] on fixture.")
        self.assertEqual(len(result.flags), 1)
        self.assertIn("Part number needed (PN?)", result.flags[0])

    def test_pn_placeholder_with_punctuation(self):
        self.assertEqual(self.polish("Seat PN?, then tight w W1/2").draft,
                         'Seat [part number needed], then tighten with a 1/2" wrench.')

    def test_nopart_opts_out_and_disappears(self):
        result = self.assert_draft("Insp for gaps NOPART", "Inspect for gaps.")
        self.assertEqual(result.status, "Ready")
        self.assert_draft("Insp for gaps. NOPART.", "Inspect for gaps.")

    def test_unknown_part_number_is_flagged_once(self):
        result = self.polish("Insert P-999 into base")
        self.assertEqual(len(result.flags), 1, result.flags)
        self.assertIn("Unknown part number: P-999", result.flags[0])

    def test_when_part_mentioned_mode(self):
        g = glossary(settings={"require_part_number": "when_part_mentioned"})
        self.assertEqual(self.polish("Insp for gaps", g).flags, [])
        self.assertIn("No part number", " ".join(self.polish("Insert the bracket", g).flags))

    def test_no_parts_sheet_still_flags(self):
        result = polish_note("Tight w W1/2", self.g, [], build_vocabulary(self.g, []))
        self.assertIn("No part number", " ".join(result.flags))


class FlagTests(EngineTestCase):
    def test_typos_are_flagged_not_guessed(self):
        result = self.polish("Tght w W1/2 P-004")
        self.assertEqual(result.draft, 'Tght with a 1/2" wrench Bracket (P-004).')
        self.assertEqual(result.unknown_words, ["Tght"])

    def test_word_forms_of_known_words_are_known(self):
        self.assertFalse(self.polish("Tightening and seated P-004 snugly").unknown_words)

    def test_part_names_and_learned_words_are_known(self):
        vocab = build_vocabulary(self.g, PARTS + [Part("P-777", "Gizmo flange")], learned=["wibble"])
        result = polish_note("Wibble the gizmo P-004", self.g, PARTS, vocab)
        self.assertEqual(result.unknown_words, [])

    def test_hard_leaves_description_blank(self):
        result = self.polish("HARD route harness thru clips")
        self.assertEqual((result.draft, result.status), ("", "HARD"))

    def test_lowercase_hard_is_just_a_word(self):
        self.assertNotEqual(self.polish("Press hard on P-004").status, "HARD")

    def test_safety_words_force_review(self):
        result = self.polish("Torque P-004 to 25 Nm")
        self.assertEqual(result.status, "Flagged")
        self.assertIn("Safety word(s): torque", " ".join(result.flags))

    def test_empty_note(self):
        self.assertEqual(self.polish("   ").status, "Flagged")


class GlossaryTests(unittest.TestCase):
    def test_starter_glossary_is_valid(self):
        self.assertGreater(len(build_glossary(STARTER_GLOSSARY).terms), 20)

    def assert_rejected(self, change, message):
        data = copy.deepcopy(STARTER_GLOSSARY)
        change(data)
        with self.assertRaises(ValueError) as caught:
            build_glossary(data)
        self.assertIn(message, str(caught.exception))

    def test_bad_glossaries_are_rejected_with_a_reason(self):
        self.assert_rejected(lambda d: d["terms"].update({"w1/2": "x"}), "both a tool code and a term")
        self.assert_rejected(lambda d: d["terms"].update({"mm": "x"}), "both a term and a unit")
        self.assert_rejected(lambda d: d["tools"].update({"w1/2": {"name": "dup"}}), "listed twice")
        self.assert_rejected(lambda d: d["terms"].update({"Tight": "x"}), "listed twice")
        self.assert_rejected(lambda d: d["tools"].update({"X1": {"phrase": "no name"}}), "needs a \"name\"")
        self.assert_rejected(lambda d: d["tools"].update({"two words": {"name": "x"}}), "one word")
        self.assert_rejected(lambda d: d["settings"].update({"part_format": "{name}"}), "must contain {number}")
        self.assert_rejected(lambda d: d["settings"].update({"require_part_number": "sometimes"}), "must be one of")
        self.assert_rejected(lambda d: d["settings"].update({"part_number_pattern": "("}), "not a valid pattern")
        self.assert_rejected(lambda d: d.update({"extras": {}}), "Unknown section")
        self.assert_rejected(lambda d: d.update({"words": "fixture"}), "must be a list")

    def test_invalid_json_points_at_the_line(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "glossary.json"
            path.write_text('{\n  "terms": {"w": "with",}\n}\n', encoding="utf-8")
            with self.assertRaises(ValueError) as caught:
                wig_glossary.load_glossary(path)
            self.assertIn("line 2", str(caught.exception))
            self.assertIn("comma", str(caught.exception))

    def test_duplicate_keys_in_the_file_are_caught(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "glossary.json"
            path.write_text('{"terms": {"w": "with", "w": "width"}}', encoding="utf-8")
            with self.assertRaises(ValueError) as caught:
                wig_glossary.load_glossary(path)
            self.assertIn("appears twice", str(caught.exception))

    def test_add_commands_edit_the_file_and_validate_first(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "glossary.json"
            wig_glossary.add_term(path, "Snug", "snug down")      # creates the starter first
            wig_glossary.add_tool(path, "W12", "12 mm wrench", "a 12 mm wrench")
            g = wig_glossary.add_words(path, ["Gizmo", "gizmo"])
            self.assertEqual(g.terms["snug"], "snug down")
            self.assertEqual(g.tools["w12"].phrase, "a 12 mm wrench")
            self.assertIn("gizmo", g.words)
            self.assertEqual(json.loads(path.read_text(encoding="utf-8"))["words"].count("gizmo"), 1)
            before = path.read_text(encoding="utf-8")
            with self.assertRaises(ValueError):
                wig_glossary.add_term(path, "w12", "clash with the tool")   # would make the file invalid
            with self.assertRaises(ValueError):
                wig_glossary.add_term(path, "snug", "something else")
            self.assertEqual(path.read_text(encoding="utf-8"), before)      # untouched on failure

    def test_starter_is_never_overwritten(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "glossary.json"
            path.write_text('{"terms": {}}', encoding="utf-8")
            self.assertFalse(wig_glossary.write_starter_glossary(path))
            self.assertEqual(path.read_text(encoding="utf-8"), '{"terms": {}}')


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)
        self.workbook = self.root / "12-345-67.xlsx"
        make_workbook(self.workbook)
        wb = load_workbook(self.workbook)
        for number, _bom, qty, name in [("P-004", "PURCHASED", 1, "Bracket"), ("P-001", "PURCHASED", 2, "Dowel pin")]:
            wb["Parts"].append([number, _bom, qty, name])
        for i, note in enumerate(["Tight w W1/2 P-004", "Insp for gaps NOPART", "Seat PN? on fixture", "HARD route it",
                                  "Tght w W1/2 P-004"], start=1):
            wb["Assembly"].append([i, note, note])   # what the OneNote import writes
        wb["Testing"].append([1, "Chk fit P-001", "Chk fit P-001"])
        wb.save(self.workbook)
        self.g = glossary()

    def run_polish(self):
        return wig_polish.polish_workbook(self.workbook, self.g, learned_words_path=self.root / "learned-words.txt",
                                          to_add_path=self.root / "to-add.csv")

    def rows(self, sheet="Assembly"):
        ws = load_workbook(self.workbook)[sheet]
        return [[c.value for c in r][:7] for r in ws.iter_rows(min_row=2)]

    def test_statuses_drafts_and_report(self):
        report = self.run_polish()
        rows = self.rows()
        self.assertEqual([r[3] for r in rows], ["Ready", "Ready", "Flagged", "HARD", "Flagged"])
        self.assertEqual(rows[0][1], 'Tighten with a 1/2" wrench Bracket (P-004).')
        self.assertIsNone(rows[3][1])                     # HARD: blank for you to write
        self.assertEqual(rows[0][2], "Tight w W1/2 P-004")  # raw note kept
        self.assertEqual(report.counts, {"Ready": 3, "Flagged": 2, "HARD": 1, "Approved": 0})
        self.assertEqual(self.rows("Testing")[0][1], "Check fit Dowel pin (P-001).")
        self.assertEqual(report.sections, ["Assembly", "Testing"])

    def test_to_resolve_sheet_lists_what_needs_you(self):
        self.run_polish()
        ws = load_workbook(self.workbook)["To Resolve"]
        listed = [(r[0], r[1]) for r in ws.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(listed, [("Assembly", 3), ("Assembly", 4), ("Assembly", 5)])
        self.assertEqual(listed, [(a, int(b)) for a, b in listed])   # numbers, not text

    def test_approved_rows_are_left_alone(self):
        self.run_polish()
        wb = load_workbook(self.workbook)
        wb["Assembly"].cell(row=2, column=2).value = "Tighten the bracket with a 1/2\" wrench."
        wb["Assembly"].cell(row=2, column=4).value = "Approved"
        wb.save(self.workbook)
        report = self.run_polish()
        self.assertEqual(report.counts["Approved"], 1)
        self.assertEqual(self.rows()[0][1], "Tighten the bracket with a 1/2\" wrench.")
        self.assertEqual(self.rows()[0][3], "Approved")

    def test_changing_an_approved_note_reopens_it(self):
        self.run_polish()
        wb = load_workbook(self.workbook)
        ws = wb["Assembly"]
        ws.cell(row=2, column=4).value = "Approved"
        ws.cell(row=2, column=3).value = "Tight w W1/2 P-001"
        wb.save(self.workbook)
        report = self.run_polish()
        self.assertEqual(report.reopened, 1)
        row = self.rows()[0]
        self.assertEqual(row[3], "Flagged")
        self.assertIn("Note changed after approval", row[4])
        self.assertEqual(row[1], 'Tighten with a 1/2" wrench Dowel pin (P-001).')

    def test_hand_edits_are_never_overwritten(self):
        self.run_polish()
        wb = load_workbook(self.workbook)
        wb["Assembly"].cell(row=6, column=2).value = "Tighten the bracket by hand."   # step 5, still Flagged
        wb.save(self.workbook)
        report = self.run_polish()
        row = self.rows()[4]
        self.assertEqual(row[1], "Tighten the bracket by hand.")
        self.assertIn("edited by hand", row[4])
        self.assertEqual(report.kept_edits, 1)

    def test_glossary_fixes_apply_on_the_next_run(self):
        self.run_polish()
        self.g = wig_glossary.build_glossary({**copy.deepcopy(STARTER_GLOSSARY),
                                              "terms": {**STARTER_GLOSSARY["terms"], "tght": "tighten"}})
        self.run_polish()
        self.assertEqual(self.rows()[4][1], 'Tighten with a 1/2" wrench Bracket (P-004).')
        self.assertEqual(self.rows()[4][3], "Ready")

    def test_unknown_words_go_to_the_to_add_list_and_accumulate(self):
        self.run_polish()
        self.run_polish()
        with (self.root / "to-add.csv").open(encoding="utf-8") as file:
            rows = list(csv.DictReader(file))
        self.assertEqual([(r["word"], r["count"]) for r in rows], [("tght", "2")])
        self.assertEqual(rows[0]["example"], "Tght w W1/2 P-004")

    def test_words_from_approved_steps_are_learned(self):
        self.run_polish()
        wb = load_workbook(self.workbook)
        ws = wb["Assembly"]
        ws.cell(row=2, column=2).value = "Frobnicate the bracket."
        ws.cell(row=2, column=4).value = "Approved"
        ws.cell(row=7 - 1, column=3).value = "Frobnicate P-004"
        wb.save(self.workbook)
        report = self.run_polish()
        self.assertIn("frobnicate", (self.root / "learned-words.txt").read_text(encoding="utf-8").split())
        self.assertGreaterEqual(report.learned, 1)
        self.assertNotIn("frobnicate", report.unknown_words)

    def test_description_only_rows_are_polished_too(self):
        wb = load_workbook(self.workbook)
        wb["Packaging"].append([1, "Bag P-004"])
        wb.save(self.workbook)
        self.run_polish()
        row = self.rows("Packaging")[0]
        self.assertEqual((row[1], row[2]), ("Bag Bracket (P-004).", "Bag P-004"))

    def test_tools_column_lists_each_tool_once_in_order_of_use(self):
        wb = load_workbook(self.workbook)
        wb["Assembly"].append([6, "Tight w W1/2 then PH2, then W1/2 again P-004", "Tight w W1/2 then PH2, then W1/2 again P-004"])
        wb.save(self.workbook)
        self.run_polish()
        ws = load_workbook(self.workbook)["Assembly"]
        self.assertEqual(ws.cell(row=1, column=8).value, "Tools")
        self.assertEqual(ws.cell(row=2, column=8).value, '1/2" wrench (W1/2)')
        self.assertEqual(ws.cell(row=7, column=8).value, '1/2" wrench (W1/2); #2 Phillips screwdriver (PH2)')
        self.assertIsNone(ws.cell(row=3, column=8).value)       # no tool in "Insp for gaps NOPART"

    def test_a_backup_is_kept(self):
        report = self.run_polish()
        self.assertTrue(report.backup_path.exists())

    def test_missing_workbook_is_a_clear_error(self):
        with self.assertRaises(ValueError) as caught:
            wig_polish.polish_workbook(self.root / "nope.xlsx", self.g)
        self.assertIn("No workbook", str(caught.exception))


if __name__ == "__main__":
    unittest.main()
