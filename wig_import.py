"""Import a OneNote page, exported as a Word .docx, into a WIG assembly.

In OneNote, put each step in one row of a table: the photo in one cell, the
description in the other. Export the page (File > Export > Current: Page >
Word Document). This reads that table and writes:

    <assembly folder>/Assembly/1.jpg, 2.jpg, ...   one photo per step, in row order
    <assembly folder>/<assembly number>.xlsx       Assembly sheet: Step # | Description | Raw note

Rows pair by position, so photos and descriptions cannot drift apart. The
columns are found by content, not by position, and OneNote's empty padding
rows are skipped. Photos are rotated upright from their EXIF flag, flattened
onto white if they have transparency, and shrunk so the deck stays small.

This module takes explicit paths and imports nothing from the generator, so it
behaves the same whether the generator runs as a script or as the `wig` command.
"""

from __future__ import annotations

import io
import re
import shutil
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional, Tuple
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from PIL import Image, ImageOps, UnidentifiedImageError

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"w": W_NS, "a": A_NS}

DEFAULT_MAX_PX = 1600
IMAGE_FORMATS = {"jpg": ("JPEG", ".jpg"), "png": ("PNG", ".png")}
STEP_IMAGE_NAME = re.compile(r"\d+\.(png|jpe?g)", re.IGNORECASE)
SHEET_ASSEMBLY = "Assembly"


@dataclass
class DocRow:
    """One non-empty table row: the 1-based row number in the table, its text, its photos."""
    row_number: int
    text: str
    images: List[bytes]


@dataclass
class ImportResult:
    steps: int
    title: str
    workbook_created: bool
    backup_path: Optional[Path]
    removed_images: int
    warnings: List[str] = field(default_factory=list)


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("\xa0", " ")).strip()


def _paragraph_text(paragraph) -> str:
    parts: List[str] = []
    for node in paragraph.iter():
        tag = node.tag.rsplit("}", 1)[-1]
        if tag == "t":
            parts.append(node.text or "")
        elif tag in ("tab", "br"):
            parts.append(" ")
    return "".join(parts)


def _cell_text(cell) -> str:
    return _clean(" ".join(_paragraph_text(p) for p in cell.findall(".//w:p", NS)))


def _relationships(archive: zipfile.ZipFile) -> dict:
    try:
        root = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
    except KeyError:
        return {}
    rels = {}
    for rel in root.findall(f"{{{PKG_REL_NS}}}Relationship"):
        rels[rel.get("Id")] = (rel.get("Target", ""), rel.get("TargetMode"))
    return rels


def _image_bytes(archive: zipfile.ZipFile, rels: dict, rel_id: str) -> bytes:
    target, mode = rels.get(rel_id, ("", None))
    if not target or mode == "External":
        raise ValueError(
            "A picture in the Word file is linked from outside the document rather than "
            "embedded, so it cannot be imported. Re-export the page from OneNote."
        )
    path = target.lstrip("/") if target.startswith("/") else f"word/{target}"
    try:
        return archive.read(path)
    except KeyError:
        raise ValueError(f"The Word file refers to a picture it does not contain: {target}")


def read_onenote_docx(path: Path) -> Tuple[str, List[DocRow]]:
    """Returns (page title, non-empty rows) from the first table in the file that holds photos."""
    try:
        archive = zipfile.ZipFile(path)
        document = ET.fromstring(archive.read("word/document.xml"))
    except (zipfile.BadZipFile, KeyError, ET.ParseError, OSError) as exc:
        raise ValueError(f"{path} is not a readable Word (.docx) file: {exc}")

    body = document.find("w:body", NS)
    if body is None:
        raise ValueError(f"{path} has no document body.")
    rels = _relationships(archive)

    title = ""
    for paragraph in body.findall("w:p", NS):
        title = _clean(_paragraph_text(paragraph))
        if title:
            break

    for table in body.findall("w:tbl", NS):
        rows: List[DocRow] = []
        for row_number, tr in enumerate(table.findall("w:tr", NS), start=1):
            texts: List[str] = []
            images: List[bytes] = []
            for cell in tr.findall("w:tc", NS):
                text = _cell_text(cell)
                if text:
                    texts.append(text)
                for blip in cell.iterfind(".//a:blip", NS):
                    rel_id = blip.get(f"{{{R_NS}}}embed")
                    if rel_id:
                        images.append(_image_bytes(archive, rels, rel_id))
            text = _clean(" ".join(texts))
            if text or images:
                rows.append(DocRow(row_number, text, images))
        if any(row.images for row in rows):
            return title, rows

    raise ValueError(
        "No table with photos was found in the Word file.\n"
        "In OneNote, put each step in one row of a table (Insert > Table), with the photo "
        "in one cell and the description in the other, then export the page again."
    )


def _validate(rows: List[DocRow]) -> List[str]:
    """Checks the rows pair up. Returns warnings, raises ValueError on a real problem."""
    problems: List[str] = []
    warnings: List[str] = []
    for row in rows:
        snippet = row.text[:40] + ("..." if len(row.text) > 40 else "")
        if len(row.images) > 1:
            problems.append(f"Table row {row.row_number} has {len(row.images)} photos; a step takes one.")
        elif not row.images:
            problems.append(f'Table row {row.row_number} has a description but no photo: "{snippet}"')
        elif not row.text:
            warnings.append(f"Table row {row.row_number} has a photo but no description; the step is blank.")
    if problems:
        raise ValueError("The table does not pair up photos with descriptions:\n  " + "\n  ".join(problems))
    if not rows:
        raise ValueError("The table has no steps.")
    return warnings


def prepare_image(data: bytes, max_px: int) -> Image.Image:
    """Opens a photo, applies its EXIF rotation, flattens transparency onto white, and shrinks it."""
    try:
        image = Image.open(io.BytesIO(data))
        image.load()
    except (UnidentifiedImageError, OSError) as exc:
        raise ValueError(f"A picture in the Word file could not be read: {exc}")
    image = ImageOps.exif_transpose(image)
    if image.mode in ("RGBA", "LA") or (image.mode == "P" and "transparency" in image.info):
        rgba = image.convert("RGBA")
        flat = Image.new("RGB", rgba.size, (255, 255, 255))
        flat.paste(rgba, mask=rgba.getchannel("A"))
        image = flat
    else:
        image = image.convert("RGB")
    if max(image.size) > max_px:
        image.thumbnail((max_px, max_px), Image.LANCZOS)
    return image


def _existing_step_images(folder: Path) -> List[Path]:
    if not folder.is_dir():
        return []
    return [p for p in folder.iterdir() if p.is_file() and STEP_IMAGE_NAME.fullmatch(p.name)]


def _assembly_rows_in(ws) -> int:
    return sum(
        1 for row in ws.iter_rows(min_row=2, values_only=True)
        if any(cell not in (None, "") for cell in row)
    )


def import_onenote(
    docx_path: Path,
    assembly_dir: Path,
    workbook_path: Path,
    *,
    replace: bool = False,
    max_px: int = DEFAULT_MAX_PX,
    image_format: str = "jpg",
) -> ImportResult:
    """Writes the photos and the Assembly sheet for one assembly from a OneNote Word export.

    Nothing is written until the whole file has been read, validated and every photo
    converted. Existing steps are never overwritten unless replace is True.
    """
    if image_format not in IMAGE_FORMATS:
        raise ValueError(f"Image format must be one of: {', '.join(IMAGE_FORMATS)}.")
    if max_px < 200:
        raise ValueError("--max-px must be at least 200.")
    pil_format, extension = IMAGE_FORMATS[image_format]

    title, rows = read_onenote_docx(Path(docx_path))
    warnings = _validate(rows)
    prepared = [prepare_image(row.images[0], max_px) for row in rows]

    image_dir = Path(assembly_dir) / SHEET_ASSEMBLY
    stale = _existing_step_images(image_dir)
    workbook_path = Path(workbook_path)
    workbook_exists = workbook_path.exists()
    wb = load_workbook(workbook_path) if workbook_exists else None
    if wb is not None and SHEET_ASSEMBLY not in wb.sheetnames:
        raise ValueError(f"{workbook_path.name} has no '{SHEET_ASSEMBLY}' sheet to put the steps in.")
    existing_rows = _assembly_rows_in(wb[SHEET_ASSEMBLY]) if wb is not None else 0
    if (stale or existing_rows) and not replace:
        raise ValueError(
            f"Assembly {Path(assembly_dir).name} already has {existing_rows} step row(s) and "
            f"{len(stale)} step photo(s). Nothing was changed. Run again with --replace to overwrite them."
        )

    backup_path: Optional[Path] = None
    if wb is not None:
        backup_path = workbook_path.with_name(workbook_path.name + ".bak")
        shutil.copy2(workbook_path, backup_path)
    else:
        from wig_scaffold import make_workbook  # headers only: Title Info etc. are still yours to fill
        make_workbook(workbook_path)
        wb = load_workbook(workbook_path)

    # A leftover 1.png would shadow a new 1.jpg, because the generator tries PNG first.
    for path in stale:
        path.unlink()
    image_dir.mkdir(parents=True, exist_ok=True)
    for number, image in enumerate(prepared, start=1):
        save_kwargs = {"quality": 90, "optimize": True} if pil_format == "JPEG" else {"optimize": True}
        image.save(image_dir / f"{number}{extension}", pil_format, **save_kwargs)

    ws = wb[SHEET_ASSEMBLY]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row)
    if not ws.cell(row=1, column=3).value:
        ws.cell(row=1, column=3).value = "Raw note"
    for number, row in enumerate(rows, start=1):
        ws.append([number, row.text, row.text])
    wb.save(workbook_path)

    return ImportResult(
        steps=len(rows),
        title=title,
        workbook_created=not workbook_exists,
        backup_path=backup_path,
        removed_images=len(stale),
        warnings=warnings,
    )
