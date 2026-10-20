"""Starter files for WIG: a blank project skeleton and an optional made-up sample.

Used by `wig --init` and by examples/make_demo.py. Nothing here is real data.
The template it builds is deliberately plain: it exists to show the slide-order
contract (see the README), not to look good. Restyle it in PowerPoint, keeping
the slide order and the {placeholders}.
"""

from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

TEMPLATE_NAME = "XX-XXX-XX REV XX.pptx"
WORKBOOK_TEMPLATE_NAME = "assembly-workbook-template.xlsx"
SAMPLE_ASSEMBLY = "12-345-67"

SAMPLE_PARTS: List[Tuple[str, str, int, str]] = [
    ("P-001", "PURCHASED", 4, "M5 x 16 socket head screw"),
    ("P-002", "PURCHASED", 4, "M5 washer"),
    ("P-003", "MFG BOM", 1, "Base plate"),
    ("P-004", "MFG BOM", 1, "Bracket"),
    ("P-005", "PURCHASED", 2, "Dowel pin"),
]
SAMPLE_STEPS: List[str] = [
    "Place the base plate on the fixture and seat both dowel pins.",
    "Set the bracket on the base plate against the dowel pins.",
    "Fasten with four screws and washers. Torque to the value on the drawing.",
    "Inspect the assembly for gaps and loose hardware.",
]

HEADERS: Dict[str, List[str]] = {
    "Title Info": ["WI #", "WI Title", "REV #", "PRODUCTION CELL", "FORMER WI#", "ASSEMBLY TIME", "Parent Assemblies"],
    "Parts": ["Part #", "BOMNAME", "Quantity", "Part Name"],
    "Equipment": ["Materials", "Tools"],
    "Assembly": ["Step #", "Description"],
    "Testing": ["Step #", "Description"],
    "Packaging": ["Step #", "Description"],
}


def make_image(path: Path, label: str, color: Tuple[int, int, int]) -> None:
    """Writes a flat-colour placeholder picture with a text label."""
    img = Image.new("RGB", (400, 300), color)
    ImageDraw.Draw(img).text((20, 20), label, fill=(255, 255, 255))
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def make_template(path: Path) -> None:
    """Builds a bare template that follows the slide-order contract in the README."""
    prs = Presentation()
    blank = prs.slide_layouts[6]
    slides = [prs.slides.add_slide(blank) for _ in range(10)]

    def text(slide, content: str, left: float, top: float, width: float = 8.5, height: float = 0.8) -> None:
        slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)) \
            .text_frame.text = content

    # The cover picture is placed on the right half, so keep cover text to the left.
    text(slides[0], "{TITLE}", 0.5, 0.4, width=9.0)
    text(slides[0], "{Assembly Number}  REV {REV}", 0.5, 1.2, width=4.3)
    text(slides[0], "Date: {DATE}", 0.5, 2.0, width=4.3)
    text(slides[0], "Cell: {CELL}", 0.5, 2.6, width=4.3)
    text(slides[0], "Assembly time: {TIME}", 0.5, 3.2, width=4.3)
    text(slides[0], "Replaces: {FORMER}", 0.5, 3.8, width=4.3)
    text(slides[0], "Used on: {PARENT_ASSEMBLIES}", 0.5, 4.4, width=4.3)
    text(slides[0], "Includes: {INCLUDED_SUBASSEMBLIES}", 0.5, 5.0, width=4.3, height=1.2)
    table = slides[1].shapes.add_table(8, 5, Inches(0.5), Inches(1), Inches(9), Inches(4)).table
    for col, head in enumerate(["DWG REV", "WI REV", "Change", "By", "Date"]):
        table.cell(0, col).text = head
    text(slides[2], "Cycle Time\nLeave page blank if no data", 0.5, 0.4)
    text(slides[3], "Welding requirements", 0.5, 0.4)
    text(slides[4], "{Part Number}", 0.5, 0.2)
    text(slides[5], "{Materials placeholder} {Tools placeholder}", 0.5, 0.2)
    text(slides[6], "{Assembly Step picture} {Step Description}", 0.5, 0.2)
    text(slides[7], "Size and Weight\nLeave page blank if no data", 0.5, 0.4)
    text(slides[8], "Testing", 0.5, 0.2)
    text(slides[9], "Packaging", 0.5, 0.2)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)


def make_workbook(path: Path, sample: bool = False) -> None:
    """Writes a workbook with the required sheets: headers only, or filled with sample data."""
    wb = Workbook()
    wb.remove(wb.active)
    for sheet, headers in HEADERS.items():
        ws = wb.create_sheet(sheet)
        ws.append(headers)
    if sample:
        wb["Title Info"].append([SAMPLE_ASSEMBLY, "Demo Bracket Assembly", "A", "Cell 1", "N/A", "12 min", "98-765-43"])
        for row in SAMPLE_PARTS:
            wb["Parts"].append(list(row))
        wb["Equipment"].append(["Thread locker", "Torque driver"])
        wb["Equipment"].append(["Clean wipes", "Assembly fixture"])
        for number, description in enumerate(SAMPLE_STEPS, start=1):
            wb["Assembly"].append([number, description])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def scaffold_project(root: Path, sample: bool = False) -> Dict[str, Path]:
    """Creates an empty WIG project under root and returns the paths it made.

    With sample=True it also fills in one made-up assembly (12-345-67) with
    images and a workbook, ready to generate.
    """
    root = Path(root)
    design = root / "Design Folder"
    catalog = root / "Part_Catalog_Images"
    template = root / "Templates" / TEMPLATE_NAME
    design.mkdir(parents=True, exist_ok=True)
    catalog.mkdir(parents=True, exist_ok=True)
    make_template(template)
    make_workbook(root / "Templates" / WORKBOOK_TEMPLATE_NAME)
    paths = {"design_folder": design, "part_images_folder": catalog, "template_path": template}

    if sample:
        assembly_dir = design / SAMPLE_ASSEMBLY.split("-")[0] / SAMPLE_ASSEMBLY
        for index, (part_number, *_rest) in enumerate(SAMPLE_PARTS):
            make_image(catalog / f"{part_number}.png", part_number, (40 + index * 30, 90, 160))
        make_image(catalog / f"{SAMPLE_ASSEMBLY}.png", f"Assembly {SAMPLE_ASSEMBLY}", (30, 120, 70))
        for number in range(1, len(SAMPLE_STEPS) + 1):
            make_image(assembly_dir / "Assembly" / f"{number}.png", f"Step {number}", (150, 70, 40))
        make_workbook(assembly_dir / f"{SAMPLE_ASSEMBLY}.xlsx", sample=True)
        paths["assembly_dir"] = assembly_dir
    return paths
